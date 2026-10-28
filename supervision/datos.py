# -*- coding: utf-8 -*-
"""
DATOS DEL PANEL DE SUPERVISION (sin Streamlit, para poder probarlo)
===================================================================
Junta en una sola lista el trabajo guardado en la base de datos (revisiones de SSAA, altas y
cambios de fichas de contrato, solicitudes de ofertas y su seguimiento, estudios de potencia,
consultas de la descarga de datos de la API de Gemweb),
lo resume por usuario y mes, y lo exporta a Excel junto con la actividad (accesos, paginas,
consultas a Gemweb y descargas, tabla `actividad` de almacen.py).

Las fechas de filtro son date o 'AAAA-MM-DD', ambas incluidas.
"""

import datetime as dt
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import almacen

SSAA, FICHAS, OFERTAS, POTENCIA, GEMWEB = ("Revisión de SSAA", "Fichas de contrato",
                                           "Solicitud de ofertas", "Estudio de potencia",
                                           "Descarga datos API Gemweb")
MODULOS = [SSAA, FICHAS, OFERTAS, POTENCIA, GEMWEB]
SIN_USUARIO = "(sin usuario)"          # fichas importadas del JSON inicial
INDICADORES = {SSAA: "Revisiones de SSAA", FICHAS: "Fichas de contrato",
               OFERTAS: "Solicitudes de ofertas", POTENCIA: "Estudios de potencia",
               GEMWEB: "Descargas API Gemweb"}
CONSULTA, DESCARGA, ACCESO = "Consulta Gemweb", "Descarga", "Inicio de sesión"
PERIODOS = ["Este mes", "Mes anterior", "Últimos 3 meses", "Este año", "Personalizado"]
LIMITE = 1_000_000                  # sin limite practico: se filtra por fecha despues

GRANATE, BLANCO = "970000", "FFFFFF"


def _iso(fecha):
    return fecha.isoformat() if isinstance(fecha, dt.date) else str(fecha)


def _dentro(fecha, desde, hasta):
    return bool(fecha) and _iso(desde) <= fecha[:10] <= _iso(hasta)


def periodo(nombre, hoy=None):
    """(desde, hasta) de un periodo predefinido; None en «Personalizado»."""
    hoy = hoy or dt.date.today()
    inicio_mes = hoy.replace(day=1)
    if nombre == "Este mes":
        return inicio_mes, hoy
    if nombre == "Mes anterior":
        fin = inicio_mes - dt.timedelta(days=1)
        return fin.replace(day=1), fin
    if nombre == "Últimos 3 meses":
        m = inicio_mes.month - 2
        return inicio_mes.replace(year=inicio_mes.year + (m - 1) // 12,
                                  month=(m - 1) % 12 + 1), hoy
    if nombre == "Este año":
        return hoy.replace(month=1, day=1), hoy
    return None


def periodo_anterior(desde, hasta):
    """Periodo de la misma duracion justo antes, para comparar."""
    dias = (hasta - desde).days + 1
    return desde - dt.timedelta(days=dias), desde - dt.timedelta(days=1)


def trabajo(desde, hasta, usuarios=None, modulos=None):
    """Trabajo guardado entre dos fechas, del mas reciente al mas antiguo. Cada registro:
    fecha (ISO), usuario, modulo, tipo, descripcion, cups, resultado, origen e id (para abrirlo)."""
    filas = []
    for r in almacen.listar_revisiones(limite=LIMITE):
        if _dentro(r["fecha"], desde, hasta):
            dif = r["diferencia"]
            filas.append(dict(
                fecha=r["fecha"], usuario=r["usuario"], modulo=SSAA, tipo="Revisión de factura",
                descripcion=" · ".join(x for x in (r["comercializadora"], r["factura"]) if x),
                cups=r["cups"], resultado="%s (%s €)" % (r["veredicto"] or "—", _num(dif))
                if dif is not None else r["veredicto"], origen="revision", id=r["id"]))
    for r in almacen.listar_historial_contratos(_iso(desde), _iso(hasta)):
        filas.append(dict(
            fecha=r["fecha"], usuario=r["usuario"], modulo=FICHAS,
            tipo="Ficha %s" % r["accion"], descripcion="Ficha de contrato", cups=r["cups"],
            resultado=r["accion"].capitalize(), origen="ficha", id=r["id"]))
    for r in almacen.listar_solicitudes(limite=LIMITE):
        producto = "gas" if r.get("producto") == "gas" else "electricidad"
        descripcion = " · ".join(x for x in (r["ref"], r["cliente"]) if x)
        if _dentro(r["fecha"], desde, hasta):
            filas.append(dict(
                fecha=r["fecha"], usuario=r["usuario"], modulo=OFERTAS,
                tipo="Solicitud de %s" % producto, descripcion=descripcion, cups=r["cups"],
                resultado="%s · %s suministro%s" % (r["estado"], r["suministros"] or 0,
                                                    "" if r["suministros"] == 1 else "s"),
                origen="solicitud", id=r["id"]))
        if r["actualizado"] and _dentro(r["actualizado"], desde, hasta):
            filas.append(dict(
                fecha=r["actualizado"], usuario=r["actualizado_por"], modulo=OFERTAS,
                tipo="Seguimiento de solicitud", descripcion=descripcion, cups=r["cups"],
                resultado=r["estado"], origen="solicitud", id=r["id"]))
    for r in almacen.listar_estudios_potencia(limite=LIMITE):
        if _dentro(r["fecha"], desde, hasta):
            filas.append(dict(
                fecha=r["fecha"], usuario=r["usuario"], modulo=POTENCIA,
                tipo="Estudio %s" % (r["tipo"] or "").lower(), descripcion=r["titular"] or "",
                cups=r["cups"], resultado="Ahorro %s €/año · inversión %s €" % (
                    _num(r["ahorro"]), _num(r["inversion"])), origen="estudio", id=r["id"]))
    for e in almacen.listar_actividad(_iso(desde), _iso(hasta), limite=LIMITE):
        if e["accion"] == CONSULTA:
            filas.append(dict(
                fecha=e["fecha"], usuario=e["usuario"], modulo=GEMWEB, tipo="Consulta a la API",
                descripcion=e["detalle"] or "", cups=None, resultado=None, origen="gemweb",
                id=e["id"]))
    if usuarios:
        filas = [f for f in filas if f["usuario"] in usuarios]
    if modulos:
        filas = [f for f in filas if f["modulo"] in modulos]
    return sorted(filas, key=lambda f: f["fecha"], reverse=True)


def actividad(desde, hasta, usuarios=None):
    """Actividad del periodo. Los intentos de acceso con un usuario que no existe quedan sin
    usuario (con el nombre tecleado en el detalle): no son nadie del equipo."""
    existen = {u["usuario"] for u in almacen.listar_usuarios()}
    filas = []
    for f in almacen.listar_actividad(_iso(desde), _iso(hasta), limite=LIMITE):
        if f["usuario"] and f["usuario"] not in existen:
            f = dict(f, usuario=None, detalle="Usuario «%s»: %s" % (f["usuario"],
                                                                    f["detalle"] or ""))
        filas.append(f)
    return [f for f in filas if f["usuario"] in usuarios] if usuarios else filas


def _num(v, d=2):
    if v is None:
        return "—"
    return ("{:,.%df}" % d).format(v).replace(",", "X").replace(".", ",").replace("X", ".")


def nombre(usuario, nombres):
    """Nombre del usuario; SIN_USUARIO en lo que no tiene (fichas importadas)."""
    return nombres.get(usuario, usuario or SIN_USUARIO)


def nombres_usuarios():
    """{usuario: nombre} de todos, tambien los dados de baja."""
    return {u["usuario"]: u["nombre"] for u in almacen.listar_usuarios()}


def resumen_por_usuario(registros, eventos, nombres, ultimos_accesos=None):
    """Una fila por usuario con trabajo o actividad: trabajos por modulo, total, consultas a
    Gemweb, descargas, inicios de sesion, dias con actividad y ultimo acceso."""
    ultimos_accesos = ultimos_accesos or {}
    quienes = sorted(({f["usuario"] for f in registros} | {e["usuario"] for e in eventos})
                     - {None}, key=lambda u: nombres.get(u, u).lower())
    filas = []
    for u in quienes:
        mios = [f for f in registros if f["usuario"] == u]
        suyos = [e for e in eventos if e["usuario"] == u]
        fila = {"Usuario": nombre(u, nombres)}
        for m in MODULOS:
            fila[m] = sum(f["modulo"] == m for f in mios)
        fila["Total trabajos"] = len(mios)
        fila["Descargas"] = sum(e["accion"] == DESCARGA for e in suyos)
        fila["Inicios de sesión"] = sum(e["accion"] == ACCESO for e in suyos)
        fila["Días con actividad"] = len({x["fecha"][:10] for x in mios + suyos})
        fila["Último acceso"] = ultimos_accesos.get(u)
        filas.append(fila)
    return filas


def por_mes(registros, nombres, por="usuario"):
    """DataFrame mes (AAAA-MM) x usuario (o x modulo, por="modulo") con el numero de trabajos."""
    if not registros:
        return pd.DataFrame()
    columna = "Módulo" if por == "modulo" else "Usuario"
    df = pd.DataFrame({"Mes": [f["fecha"][:7] for f in registros],
                       columna: [f["modulo"] if por == "modulo"
                                 else nombre(f["usuario"], nombres) for f in registros]})
    return df.pivot_table(index="Mes", columns=columna, aggfunc="size", fill_value=0)


def por_usuario_y_modulo(registros, nombres):
    """DataFrame usuario x modulo con el numero de trabajos."""
    if not registros:
        return pd.DataFrame()
    df = pd.DataFrame({"Usuario": [nombre(f["usuario"], nombres)
                                   for f in registros],
                       "Módulo": [f["modulo"] for f in registros]})
    return df.pivot_table(index="Usuario", columns="Módulo", aggfunc="size", fill_value=0)


# ------------------------------------------------------------------- Excel
def _hoja(wb, titulo, cabecera, filas, anchos):
    ws = wb.create_sheet(titulo)
    ws.append(cabecera)
    for c in ws[1]:
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=GRANATE)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for f in filas:
        ws.append(f)
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _fecha(txt):
    return dt.datetime.fromisoformat(txt) if txt else None


def excel_supervision(registros, eventos, nombres, desde, hasta, ultimos_accesos=None):
    """Excel con tres hojas: Resumen (por usuario), Trabajo y Actividad."""
    wb = Workbook()
    wb.remove(wb.active)
    resumen = resumen_por_usuario(registros, eventos, nombres, ultimos_accesos)
    cab = list(resumen[0]) if resumen else ["Usuario"] + MODULOS
    ws = _hoja(wb, "Resumen", cab,
               [[_fecha(v) if k == "Último acceso" else v for k, v in f.items()]
                for f in resumen], [28] + [14] * (len(cab) - 1))
    ws.insert_rows(1, 2)
    ws["A1"] = "Trabajo del equipo del %s al %s" % (
        _iso(desde)[8:10] + "/" + _iso(desde)[5:7] + "/" + _iso(desde)[:4],
        _iso(hasta)[8:10] + "/" + _iso(hasta)[5:7] + "/" + _iso(hasta)[:4])
    ws["A1"].font = Font(bold=True, size=13, color=GRANATE)
    ws.freeze_panes = "A4"
    if resumen:
        ws.auto_filter.ref = "A3:%s%d" % (get_column_letter(len(cab)), len(resumen) + 3)
        for fila in ws.iter_rows(min_row=4, min_col=len(cab), max_col=len(cab)):
            fila[0].number_format = "DD/MM/YYYY HH:MM"
    ws = _hoja(wb, "Trabajo",
               ["Fecha", "Usuario", "Módulo", "Tipo", "Descripción", "CUPS", "Resultado"],
               [[_fecha(f["fecha"]), nombre(f["usuario"], nombres), f["modulo"],
                 f["tipo"], f["descripcion"], f["cups"], f["resultado"]] for f in registros],
               [17, 24, 20, 24, 36, 26, 36])
    for c in ws["A"][1:]:
        c.number_format = "DD/MM/YYYY HH:MM"
    ws = _hoja(wb, "Actividad", ["Fecha", "Usuario", "Módulo", "Acción", "Detalle"],
               [[_fecha(e["fecha"]), nombre(e["usuario"], nombres), e["modulo"],
                 e["accion"], e["detalle"]] for e in eventos],
               [17, 24, 20, 18, 60])
    for c in ws["A"][1:]:
        c.number_format = "DD/MM/YYYY HH:MM"
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
