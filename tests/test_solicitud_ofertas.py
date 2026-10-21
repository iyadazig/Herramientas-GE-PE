# -*- coding: utf-8 -*-
"""Excel de solicitud de ofertas (datos ficticios)."""

import datetime as dt
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from gemweb_extractor import solicitud_ofertas as so


def factura(anio, mes, dia, p1=0.0, p6=0.0):
    return {"Fin factura": dt.date(anio, mes, dia), "P1 kWh": p1, "P2 kWh": 0, "P3 kWh": 0,
            "P4 kWh": 0, "P5 kWh": 0, "P6 kWh": p6}


class Datos(unittest.TestCase):
    def test_alias_cif_y_discriminacion(self):
        self.assertEqual(so.alias_de("Empresa de prueba SA_TC4"), "TC4")
        self.assertEqual(so.alias_de("Sin alias"), "")
        self.assertEqual(so.cif_limpio("a-00.000.000"), "A00000000")
        self.assertEqual((so.discriminacion("2.0TD"), so.discriminacion("6.1TD")), ("DH3", "DH6"))

    def test_perfil_por_mes_de_fin_de_factura(self):
        # diciembre facturado en dos facturas (alta a mitad de mes): se suman
        perfil, aviso = so.perfil_mensual([factura(2025, 12, 15, p1=10), factura(2025, 12, 31, p1=5),
                                           factura(2026, 1, 31, p6=7)])
        self.assertEqual(perfil[12][0], 15.0)
        self.assertEqual(perfil[1][5], 7.0)
        self.assertEqual(perfil[6], [0.0] * 6)
        self.assertEqual(aviso, "")

    def test_aviso_si_se_mezclan_anios(self):
        _, aviso = so.perfil_mensual([factura(2025, 3, 31, p1=1), factura(2026, 3, 31, p1=1)])
        self.assertIn("Mar", aviso)

    def test_nombres_de_hojas(self):
        self.assertEqual(so.nombres_de_hojas([{"alias": "TC1"}, {"alias": "",
                                                                 "cups": "ES0000000000000000XX"}]),
                         ["2. TC1", "3. 000000XX"])


class Excel(unittest.TestCase):
    def test_libro_completo(self):
        perfil, _ = so.perfil_mensual([factura(2026, m, 28, p1=100, p6=50) for m in range(1, 13)])
        s = {"grupo": "GRUPO DE PRUEBA", "alias": "TC1", "cups": "ES0000000000000000XX",
             "razon": "EMPRESA DE PRUEBA SA",
             "direccion": "Calle Falsa 1", "cif": "A00000000", "tarifa": "6.1TD",
             "fecha_fin": "31/12/2026", "potencias": [100, 100, 100, 100, 100, 200],
             "perfil": perfil}
        cab = {"ref": "R-1", "fecha": "01/01/2026", "gestor": "Persona de prueba",
               "telefono": so.TELEFONO_GEYPE, "email": "prueba@example.com", "producto": "X"}
        wb = openpyxl.load_workbook(io.BytesIO(so.excel_solicitud(cab, [s])))
        self.assertEqual(wb.sheetnames, ["Resumen", "2. TC1"])
        h = wb["2. TC1"]
        self.assertEqual((h["C6"].value, h["D8"].value, h["H9"].value, h["D10"].value),
                         ("R-1", "Persona de prueba", "prueba@example.com", "X"))
        self.assertEqual((h["E12"].value, h["D13"].value, h["D14"].value, h["H14"].value),
                         ("EMPRESA DE PRUEBA SA", "A00000000", "ES0000000000000000XX", "TC1"))
        self.assertEqual([h[c + "21"].value for c in "DEFGHI"], [100, 100, 100, 100, 100, 200])
        self.assertEqual((h["B25"].value, h["C25"].value, h["H25"].value), ("Ene / Jan", 100, 50))
        self.assertEqual(h["I25"].value, "=SUM(C25:H25)")
        self.assertEqual(h["H23"].value, "DH6")
        self.assertEqual(len(h._charts), 1)
        r = wb["Resumen"]
        self.assertEqual([r.cell(4, n).value for n in range(1, 9)],
                         ["GRUPO DE PRUEBA", "TC1", "ES0000000000000000XX", "EMPRESA DE PRUEBA SA",
                          "Calle Falsa 1", "A00000000", "6.1TD", "31/12/2026"])
        self.assertEqual((r["A2"].value, r["I2"].value, r["O2"].value),
                         ("Cliente/Grupo", "POTENCIAS A CONTRATAR (kW)", "PREVISIÓN DE CONSUMO ANUAL"))
        self.assertEqual([r.cell(4, n).value for n in range(9, 15)], [100] * 5 + [200])
        self.assertEqual(r["O4"].value, "='2. TC1'!C37")
        self.assertEqual(r["U4"].value, "='2. TC1'!I37")


if __name__ == "__main__":
    unittest.main()
