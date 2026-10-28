# -*- coding: utf-8 -*-
"""Rol supervisor, registro de actividad y panel del equipo, con una base de datos temporal y
datos ficticios."""

import datetime as dt
import io
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

import almacen
import config
from supervision import datos as sd

RAIZ = Path(__file__).resolve().parents[1]
CUPS = "ES0000000000000000XX"
HOY = dt.date.today().isoformat()


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        carpeta = Path(self.tmp.name)
        self.p = [mock.patch.object(config, "CARPETA_DATOS", carpeta),
                  mock.patch.object(config, "FICHERO_CONTRATOS", carpeta / "no_existe.json")]
        for p in self.p:
            p.start()
        almacen._iniciado.clear()

    def tearDown(self):
        for p in self.p:
            p.stop()
        almacen._iniciado.clear()
        self.tmp.cleanup()

    def poblar(self):
        almacen.crear_usuario("ana", "Ana Prueba", "clave-segura-1", admin=True)
        almacen.crear_usuario("luis", "Luis Prueba", "clave-segura-2", supervisor=True)
        almacen.crear_usuario("eva", "Eva Prueba", "clave-segura-3")
        almacen.guardar_revision("eva", {"cups": CUPS, "comercializadora": "Comer",
                                         "factura": "F-1"},
                                 [{"Facturado €": 10.0, "Recalculado €": 8.0,
                                   "Veredicto": "FACTURADO DE MÁS"}], b"xlsx")
        almacen.guardar_contrato(CUPS, {"cups": CUPS, "techo": 1.0}, "eva")
        almacen.guardar_contrato(CUPS, {"cups": CUPS, "techo": 2.0}, "ana")
        id_sol = almacen.guardar_solicitud("ana", {"ref": "R-1", "cliente": "Cliente ficticio",
                                                   "cups": [CUPS]}, {"cabecera": {}}, "texto",
                                           b"excel")
        almacen.actualizar_seguimiento(id_sol, "Enviada", "", "eva")
        almacen.guardar_estudio_potencia("eva", {"tipo": "Individual", "titular": "Titular",
                                                 "cups": [CUPS], "ahorro": 100.0,
                                                 "inversion": 0.0}, {}, b"", b"%PDF")


class Rol(Base):
    def test_migra_columna_supervisor(self):
        ruta = almacen.ruta_bd()
        ruta.parent.mkdir(parents=True, exist_ok=True)
        con = sqlite3.connect(ruta)            # base antigua: sin supervisor ni actividad
        con.execute("CREATE TABLE usuarios (usuario TEXT PRIMARY KEY, nombre TEXT NOT NULL, "
                    "sal TEXT NOT NULL, clave TEXT NOT NULL, admin INTEGER NOT NULL DEFAULT 0, "
                    "activo INTEGER NOT NULL DEFAULT 1, cambiar_clave INTEGER NOT NULL DEFAULT 0, "
                    "creado TEXT NOT NULL, ultimo_acceso TEXT, intentos_fallidos INTEGER NOT "
                    "NULL DEFAULT 0, bloqueado_hasta TEXT)")
        con.commit()
        con.close()
        almacen.crear_usuario("eva", "Eva", "clave-segura-3")
        self.assertEqual(almacen.listar_usuarios()[0]["supervisor"], 0)
        almacen.actualizar_usuario("eva", supervisor=True)
        u = almacen.comprobar("eva", "clave-segura-3")
        self.assertTrue(almacen.puede_supervisar(u))
        self.assertFalse(u["admin"])

    def test_puede_supervisar(self):
        self.assertTrue(almacen.puede_supervisar({"admin": 1, "supervisor": 0}))
        self.assertTrue(almacen.puede_supervisar({"admin": 0, "supervisor": 1}))
        self.assertFalse(almacen.puede_supervisar({"admin": 0}))
        self.assertFalse(almacen.puede_supervisar(None))


class Actividad(Base):
    def test_accesos_quedan_registrados(self):
        almacen.crear_usuario("eva", "Eva", "clave-segura-3")
        almacen.comprobar("eva", "clave-segura-3")
        with self.assertRaises(ValueError):
            almacen.comprobar("eva", "mal-clave-99")
        acciones = [e["accion"] for e in almacen.listar_actividad()]
        self.assertEqual(acciones, ["Acceso fallido", "Inicio de sesión"])

    def test_filtra_por_fecha_y_usuario(self):
        almacen.registrar_actividad("eva", "API de Gemweb", "Consulta Gemweb", "Inventario")
        almacen.registrar_actividad("ana", "Revisión de SSAA", "Descarga", "x.xlsx")
        self.assertEqual(len(almacen.listar_actividad(HOY, HOY)), 2)
        self.assertEqual(len(almacen.listar_actividad("2000-01-01", "2000-12-31")), 0)
        self.assertEqual([e["detalle"] for e in almacen.listar_actividad(usuario="eva")],
                         ["Inventario"])

    def test_boton_descarga_anota_la_descarga(self):
        import estilo
        contenedor = mock.Mock()
        with mock.patch.object(estilo.st, "session_state", {"usuario": {"usuario": "eva"}}):
            estilo.boton_descarga("Solicitud de ofertas", contenedor, "Descargar", b"x",
                                  "sol.xlsx", "mime", type="primary")
        args, kw = contenedor.download_button.call_args
        self.assertEqual(args, ("Descargar", b"x", "sol.xlsx", "mime"))
        self.assertEqual(kw["type"], "primary")
        kw["on_click"](*kw["args"])                       # lo que hace Streamlit al pulsar
        e = almacen.listar_actividad()[0]
        self.assertEqual((e["usuario"], e["modulo"], e["accion"], e["detalle"]),
                         ("eva", "Solicitud de ofertas", "Descarga", "sol.xlsx"))

    def test_registrar_nunca_rompe(self):
        with mock.patch.object(almacen, "_con", side_effect=sqlite3.OperationalError("x")):
            almacen.registrar_actividad("eva", "m", "a")          # sin excepcion


class Datos(Base):
    def test_periodos(self):
        hoy = dt.date(2026, 1, 15)
        self.assertEqual(sd.periodo("Este mes", hoy), (dt.date(2026, 1, 1), hoy))
        self.assertEqual(sd.periodo("Mes anterior", hoy),
                         (dt.date(2025, 12, 1), dt.date(2025, 12, 31)))
        self.assertEqual(sd.periodo("Últimos 3 meses", hoy), (dt.date(2025, 11, 1), hoy))
        self.assertIsNone(sd.periodo("Personalizado", hoy))
        self.assertEqual(sd.periodo_anterior(dt.date(2026, 1, 1), dt.date(2026, 1, 31)),
                         (dt.date(2025, 12, 1), dt.date(2025, 12, 31)))

    def test_trabajo_junta_los_cuatro_modulos(self):
        self.poblar()
        registros = sd.trabajo(HOY, HOY)
        tipos = sorted(f["tipo"] for f in registros)
        self.assertEqual(tipos, ["Estudio individual", "Ficha creada", "Ficha modificada",
                                 "Revisión de factura", "Seguimiento de solicitud",
                                 "Solicitud de electricidad"])
        self.assertEqual(len(sd.trabajo(HOY, HOY, usuarios=["ana"])), 2)
        self.assertEqual(len(sd.trabajo(HOY, HOY, modulos=[sd.OFERTAS])), 2)
        self.assertEqual(sd.trabajo("2000-01-01", "2000-01-31"), [])

    def test_resumen_por_usuario(self):
        self.poblar()
        almacen.registrar_actividad("eva", "API de Gemweb", sd.CONSULTA, "Inventario")
        almacen.registrar_actividad("eva", "API de Gemweb", sd.DESCARGA, "x.xlsx")
        nombres = sd.nombres_usuarios()
        importada = dict(fecha=HOY, usuario=None, modulo=sd.FICHAS)   # fichas del JSON inicial
        filas = {f["Usuario"]: f for f in sd.resumen_por_usuario(
            sd.trabajo(HOY, HOY) + [importada], sd.actividad(HOY, HOY), nombres)}
        eva = filas["Eva Prueba"]
        self.assertEqual((eva[sd.SSAA], eva[sd.FICHAS], eva[sd.OFERTAS], eva[sd.POTENCIA]),
                         (1, 1, 1, 1))
        self.assertEqual((eva[sd.GEMWEB], eva["Total trabajos"], eva["Descargas"],
                          eva["Días con actividad"]), (1, 5, 1, 1))
        self.assertNotIn("Luis Prueba", filas)                  # sin trabajo ni actividad

    def test_acceso_con_usuario_que_no_existe(self):
        with self.assertRaises(ValueError):
            almacen.comprobar("nadie", "clave-cualquiera-1")
        self.assertNotIn("nadie", sd.nombres_usuarios())
        eventos = sd.actividad(HOY, HOY)
        fallido = [e for e in eventos if e["accion"] == "Acceso fallido"]
        self.assertEqual([e["usuario"] for e in fallido], [None])
        self.assertIn("nadie", fallido[0]["detalle"])
        self.assertEqual(sd.resumen_por_usuario([], eventos, sd.nombres_usuarios()), [])
        # registros antiguos anotados a nombre del usuario inexistente: tampoco cuentan
        almacen.registrar_actividad("fantasma", "Acceso", "Acceso fallido", "x")
        self.assertEqual(sd.resumen_por_usuario([], sd.actividad(HOY, HOY),
                                                sd.nombres_usuarios()), [])

    def test_excel_con_tres_hojas(self):
        self.poblar()
        nombres = sd.nombres_usuarios()
        datos = sd.excel_supervision(sd.trabajo(HOY, HOY), sd.actividad(HOY, HOY), nombres,
                                     dt.date.today(), dt.date.today())
        wb = load_workbook(io.BytesIO(datos))
        self.assertEqual(wb.sheetnames, ["Resumen", "Trabajo", "Actividad"])
        self.assertEqual(wb["Trabajo"].max_row, 7)               # cabecera + 6 trabajos
        self.assertEqual(wb["Resumen"]["A3"].value, "Usuario")


class Panel(Base):
    """El panel con AppTest: solo lo ve quien puede supervisar, y la descarga queda anotada."""

    def _app(self, usuario, pagina="pagina_panel"):
        from streamlit.testing.v1 import AppTest
        at = AppTest.from_string(
            "import sys\nsys.path.insert(0, %r)\nfrom supervision import paginas\n"
            "paginas.%s()\n" % (str(RAIZ), pagina), default_timeout=60)
        at.session_state["usuario"] = usuario
        return at.run()

    def test_usuario_normal_no_lo_ve(self):
        at = self._app({"usuario": "eva", "admin": 0, "supervisor": 0})
        self.assertIn("Solo los supervisores", at.warning[0].value)

    def test_supervisor_lo_ve_y_filtra(self):
        self.poblar()
        at = self._app({"usuario": "luis", "admin": 0, "supervisor": 1})
        self.assertFalse(at.exception)
        self.assertEqual(at.metric[0].value, "1")                # una revision
        at.multiselect(key="sup_usuarios").set_value(["ana"]).run()
        self.assertFalse(at.exception)
        self.assertEqual(at.metric[1].value, "1")                # ana: una ficha modificada

    def test_mi_trabajo_solo_lo_suyo(self):
        self.poblar()
        at = self._app({"usuario": "ana", "admin": 0, "supervisor": 0}, "pagina_mi_trabajo")
        self.assertFalse(at.exception)
        self.assertFalse(at.warning)                             # cualquiera puede verla
        self.assertEqual([m.value for m in at.metric[:4]], ["0", "1", "1", "0"])
        self.assertEqual(at.metric[4].label, "Descargas API Gemweb")
        self.assertEqual(at.metric[5].label, "Días con actividad")
        self.assertEqual([m.label for m in at.multiselect], ["Módulos"])   # sin filtro de usuarios
        trabajo = at.dataframe[1].value                          # 0: resumen, 1: trabajo
        self.assertNotIn("Usuario", trabajo.columns)
        self.assertEqual(len(trabajo), 2)                        # ficha modificada y solicitud


if __name__ == "__main__":
    unittest.main()
