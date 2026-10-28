# -*- coding: utf-8 -*-
"""Plantilla de suministros y consumos de la solicitud de ofertas (datos ficticios)."""

import io
import json
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from ofertas import correo
from ofertas import plantilla as pl
from ofertas import solicitud as so
from tests.test_ofertas import CAB_GAS, suministro_gas

CUPS1, CUPS2, NUEVO = "ES0000000000000001XX", "ES0000000000000002XX", "ES0000000000000009XX"
CONSIDERACIONES = {"fijo": True, "indexado": ["PVPC"], "inicio": "01/01/2027",
                   "duracion": "31/12/2027", "duracion2": "", "duracion3": "",
                   "temporada": False, "exencion": False, "gdo": False, "excedentes": False}
CAB = {"ref": "R-1", "fecha": "01/01/2026", "gestor": "Persona de prueba",
       "telefono": so.TELEFONO_GEYPE, "email": "prueba@example.com",
       "producto": so.PRODUCTOS[so.ELECTRICIDAD], "producto_tipo": so.ELECTRICIDAD,
       "recepcion_desde": "22/10/2026", "recepcion_hasta": "23/10/2026",
       "recepcion_hora": "12:00"}


def ficha(cups, p1=100.0):
    """Ficha como las de Gemweb (gemweb_extractor.paginas.ficha_solicitud)."""
    return {"cups": cups, "alias": "TC" + cups[-3], "tipo": so.ELECTRICIDAD,
            "grupo": "GRUPO DE PRUEBA", "razon": "EMPRESA DE PRUEBA SA", "cif": "A00000000",
            "direccion": "Calle Falsa 1", "tarifa": "6.1TD", "fecha_inicio": "",
            "potencias": [100, 100, 100, 100, 100, 200],
            "perfil": {m: [p1, 0, 0, 0, 0, 50.0] for m in range(1, 13)}, "aviso": "",
            "aviso_grave": False, "facturas": 12, "ultima_factura": None}


def editar(contenido, cambios):
    """Abre la plantilla, aplica cambios(wb) y la devuelve como bytes."""
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    cambios(wb)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def fila_vacia(ws):
    return next(r for r in range(2, ws.max_row + 1) if not ws.cell(r, 1).value)


def rellenar_nuevo(wb, p1=10.0, tarifa="3.0TD", con_consumos=True, gas=False):
    s = wb[pl.HOJA_SUMINISTROS]
    r = fila_vacia(s)
    valores = [NUEVO, "OTRO GRUPO", "EMPRESA NUEVA SL", "B00000000", "Nave", "Calle Inventada 2",
               tarifa] + ([300] if gas else [15, 15, 15, 15, 15, 20])
    for n, v in enumerate(valores, 1):
        s.cell(r, n, v)
    if con_consumos:
        c = wb[pl.HOJA_CONSUMOS]
        for m in range(1, 13):
            c.append([NUEVO, m, None, p1] + ([] if gas else [0, 0, 0, 0, 5]))


def cambiar_consumo(wb, cups, mes, valor, columna=4):
    c = wb[pl.HOJA_CONSUMOS]
    fila = next(r for r in range(2, c.max_row + 1)
                if c.cell(r, 1).value == cups and c.cell(r, 2).value == mes)
    c.cell(fila, columna, valor)


class Plantilla(unittest.TestCase):
    def test_sale_prellenada(self):
        wb = openpyxl.load_workbook(io.BytesIO(pl.plantilla_excel(
            [ficha(CUPS1)], [CUPS2], so.ELECTRICIDAD)))
        self.assertEqual(wb.sheetnames, [pl.HOJA_INSTRUCCIONES, pl.HOJA_SUMINISTROS,
                                         pl.HOJA_CONSUMOS])
        s, c = wb[pl.HOJA_SUMINISTROS], wb[pl.HOJA_CONSUMOS]
        self.assertEqual([s.cell(2, n).value for n in (1, 3, 7, 8, 13, 14)],
                         [CUPS1, "EMPRESA DE PRUEBA SA", "6.1TD", 100, 200, pl.GEMWEB])
        self.assertEqual((s.cell(3, 1).value, s.cell(3, 14).value), (CUPS2, pl.MANUAL))
        self.assertEqual(c.max_row, 1 + 24)                  # 12 meses x 2 CUPS
        self.assertEqual([c.cell(2, n).value for n in range(1, 5)], [CUPS1, 1, "Enero", 100])

    def test_sin_cambios_no_cambia_nada(self):
        contenido = pl.plantilla_excel([ficha(CUPS1)], [CUPS2], so.ELECTRICIDAD)
        leido = pl.leer_plantilla(contenido, so.ELECTRICIDAD)
        self.assertEqual(leido[2], [])                       # el CUPS no encontrado se ignora
        self.assertNotIn(CUPS2, leido[0])
        nuevas, cambios = pl.aplicar_plantilla([ficha(CUPS1)], leido)
        self.assertEqual(len(nuevas), 1)
        self.assertEqual(cambios[0]["Diferencia (%)"], 0)

    def test_cambia_consumo_y_añade_un_manual(self):
        fichas = [ficha(CUPS1), ficha(CUPS2)]
        contenido = editar(pl.plantilla_excel(fichas, [], so.ELECTRICIDAD), lambda wb: (
            cambiar_consumo(wb, CUPS1, 1, "1.200,5"), rellenar_nuevo(wb)))
        leido = pl.leer_plantilla(contenido, so.ELECTRICIDAD)
        self.assertEqual(leido[2] + pl.errores_de_manuales(leido[0], leido[1], fichas, False),
                         [])
        nuevas, cambios = pl.aplicar_plantilla(fichas, leido)
        uno, dos, nuevo = nuevas
        self.assertEqual(uno["perfil"][1][0], 1200.5)
        self.assertEqual(uno["origen_consumo"], pl.PLANTILLA)
        self.assertEqual(so.consumo_anual(dos), so.consumo_anual(ficha(CUPS2)))
        self.assertEqual((nuevo["cups"], nuevo["origen"], nuevo["tarifa"], nuevo["potencias"]),
                         (NUEVO, pl.MANUAL, "3.0TD", [15, 15, 15, 15, 15, 20]))
        self.assertEqual(so.consumo_anual(nuevo), 12 * 15)
        self.assertEqual([c["CUPS"] for c in cambios], [CUPS1, CUPS2, NUEVO])
        self.assertIsNone(cambios[2]["Consumo Gemweb (kWh)"])
        self.assertEqual(pl.origen_consumo(dos), pl.GEMWEB)

        # la solicitud sale con los tres: Excel con nota en los de plantilla y correo
        suministros = [dict(f, consideraciones=CONSIDERACIONES) for f in nuevas]
        wb = openpyxl.load_workbook(io.BytesIO(so.excel_solicitud(CAB, suministros)))
        self.assertEqual(len(wb.sheetnames), 4)
        hoja_nuevo = wb[wb.sheetnames[3]]
        self.assertEqual((hoja_nuevo["D14"].value, hoja_nuevo["C25"].value),
                         (NUEVO, 10))                       # P1 del mes 1
        self.assertEqual(hoja_nuevo["B39"].value, so.NOTA_CONSUMO)
        self.assertIsNone(wb[wb.sheetnames[2]]["B39"].value)      # CUPS2: de Gemweb
        total = sum(so.consumo_anual(f) for f in nuevas)
        self.assertIn("consumo anual previsto total de %s kWh"
                      % f"{total:,.0f}".replace(",", "."),
                      correo.texto_correo(CAB, suministros))

    def test_solicitud_solo_con_plantilla(self):
        contenido = editar(pl.plantilla_excel([], [], so.ELECTRICIDAD), rellenar_nuevo)
        leido = pl.leer_plantilla(contenido, so.ELECTRICIDAD)
        self.assertEqual(leido[2] + pl.errores_de_manuales(leido[0], leido[1], [], False), [])
        nuevas, _ = pl.aplicar_plantilla([], leido)
        self.assertEqual([f["cups"] for f in nuevas], [NUEVO])
        # se guarda y se vuelve a cargar con su origen (JSON, como almacen)
        vuelta = json.loads(json.dumps(nuevas[0], default=str))
        self.assertEqual((vuelta["origen"], vuelta["origen_consumo"]), (pl.MANUAL, pl.PLANTILLA))

    def test_errores(self):
        fichas = [ficha(CUPS1)]

        def estropear(wb):
            rellenar_nuevo(wb, tarifa="", con_consumos=False)
            c = wb[pl.HOJA_CONSUMOS]
            c.delete_rows(next(r for r in range(2, c.max_row + 1)
                               if c.cell(r, 1).value == CUPS1 and c.cell(r, 2).value == 12))
            cambiar_consumo(wb, CUPS1, 3, "mucho")
            c.append(["ES0000000000000005XX", 1, None, 10, 0, 0, 0, 0, 0])
        leido = pl.leer_plantilla(editar(pl.plantilla_excel(fichas, [], so.ELECTRICIDAD),
                                         estropear), so.ELECTRICIDAD)
        errores = leido[2] + pl.errores_de_manuales(leido[0], leido[1], fichas, False)
        texto = "\n".join(errores)
        self.assertIn("no es un número", texto)
        self.assertIn("%s: faltan los meses 3, 12" % CUPS1, texto)
        self.assertIn("ES0000000000000005XX: no está en la pestaña «Suministros»", texto)
        self.assertIn("%s (no está en Gemweb): falta tarifa, consumos" % NUEVO, texto)
        self.assertEqual(pl.leer_plantilla(b"no es un excel")[2][0][:20], "No se puede abrir el")

    def test_gas(self):
        g = suministro_gas(CUPS1)
        contenido = editar(pl.plantilla_excel([g], [], so.GAS), lambda wb: (
            cambiar_consumo(wb, CUPS1, 1, 1_000_000), rellenar_nuevo(
                wb, p1=500_000, tarifa="RL.6", gas=True)))
        leido = pl.leer_plantilla(contenido, so.GAS)
        self.assertEqual(leido[2] + pl.errores_de_manuales(leido[0], leido[1], [g], True), [])
        (uno, nuevo), _ = pl.aplicar_plantilla([g], leido, so.GAS)
        self.assertEqual(uno["perfil_gas"][1]["kwh"], 1_000_000)
        self.assertEqual(so.consumo_anual(uno), 1_000_000 + 11 * 100_000)
        self.assertEqual(so.consumo_anual(nuevo), 6_000_000)
        # mas de 5 GWh: telemedida y termino fijo variable, con el Qd de la plantilla
        self.assertEqual((nuevo["tecnicos"]["qd"], nuevo["tecnicos"]["telemedida"],
                          nuevo["tecnicos"]["termino_fijo"]), (300, True, so.VARIABLE))
        suministros = [dict(f, consideraciones=g["consideraciones"]) for f in (uno, nuevo)]
        wb = openpyxl.load_workbook(io.BytesIO(so.excel_solicitud(CAB_GAS, suministros)))
        self.assertEqual(wb[wb.sheetnames[2]]["B41"].value, so.NOTA_CONSUMO)


if __name__ == "__main__":
    unittest.main()
