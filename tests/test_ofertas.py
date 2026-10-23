# -*- coding: utf-8 -*-
"""Excel de solicitud de ofertas (datos ficticios)."""

import datetime as dt
import io
import json
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl
import pandas as pd

from ofertas import correo
from ofertas import solicitud as so


def factura(anio, mes, dia, p1=0.0, p6=0.0):
    return {"Fin factura": dt.date(anio, mes, dia), "P1 kWh": p1, "P2 kWh": 0, "P3 kWh": 0,
            "P4 kWh": 0, "P5 kWh": 0, "P6 kWh": p6}


class Datos(unittest.TestCase):
    def test_alias_y_cif(self):
        self.assertEqual(so.alias_de("Empresa de prueba SA_TC4"), "TC4")
        self.assertEqual(so.alias_de("Empresa de prueba_Nave_2 Norte"), "Nave_2 Norte")
        self.assertEqual(so.alias_de("Nombre sin guion bajo"), "Nombre sin guion bajo")
        self.assertEqual(so.alias_de(None), "")
        self.assertEqual(so.cif_limpio("a-00.000.000"), "A00000000")

    def test_perfil_por_mes_de_fin_de_factura(self):
        # diciembre facturado en dos facturas (alta a mitad de mes): se suman
        perfil, aviso = so.perfil_mensual([factura(2025, 12, 15, p1=10), factura(2025, 12, 31, p1=5),
                                           factura(2026, 1, 31, p6=7)])
        self.assertEqual(perfil[12][0], 15.0)
        self.assertEqual(perfil[1][5], 7.0)
        self.assertEqual(perfil[6], [0.0] * 6)
        self.assertEqual(aviso, "")

    def test_aviso_si_se_mezclan_anios(self):
        _, aviso = so.perfil_mensual([factura(2025, 3, 31, p1=1), factura(2026, 3, 31, p1=1)])
        self.assertIn("marzo", aviso)
        # mas de un año de facturas: el mismo mes se suma dos veces (grave)
        f = [dict(factura(2025, 3, 31), **{"Inicio factura": dt.date(2025, 3, 1)}),
             dict(factura(2026, 3, 31), **{"Inicio factura": dt.date(2026, 3, 1)})]
        texto, grave = so.aviso_perfil(f)
        self.assertTrue(grave)
        self.assertIn("396 días, más de un año, y en marzo", texto)
        # factura corta por un cambio de peaje: 25/09/2025-24/09/2026, sin dias repetidos
        f = [dict(factura(2025, 9, 30), **{"Inicio factura": dt.date(2025, 9, 25)}),
             dict(factura(2026, 9, 24), **{"Inicio factura": dt.date(2026, 8, 25)})]
        texto, grave = so.aviso_perfil(f)
        self.assertFalse(grave)
        self.assertTrue(texto.startswith("la fila de septiembre (2025 y 2026) junta"))
        self.assertIn("365 días", texto)
        self.assertEqual(so.aviso_perfil([factura(2026, 1, 31)]), ("", False))

    def test_nombres_de_hojas(self):
        self.assertEqual(so.nombres_de_hojas([{"alias": "TC1"}, {"alias": "",
                                                                 "cups": "ES0000000000000000XX"}]),
                         ["1. TC1", "2. 000000XX"])


class Excel(unittest.TestCase):
    def test_libro_completo(self):
        perfil, _ = so.perfil_mensual([factura(2026, m, 28, p1=100, p6=50) for m in range(1, 13)])
        s = {"grupo": "GRUPO DE PRUEBA", "alias": "TC1", "cups": "ES0000000000000000XX",
             "razon": "EMPRESA DE PRUEBA SA",
             "direccion": "Calle Falsa 1", "cif": "A00000000", "tarifa": "6.1TD",
             "fecha_inicio": "31/12/2026", "potencias": [100, 100, 100, 100, 100, 200],
             "perfil": perfil}
        cab = {"ref": "R-1", "fecha": "01/01/2026", "gestor": "Persona de prueba",
               "telefono": so.TELEFONO_GEYPE, "email": "prueba@example.com", "producto": "X"}
        wb = openpyxl.load_workbook(io.BytesIO(so.excel_solicitud(cab, [s])))
        self.assertEqual(wb.sheetnames, ["Resumen", "1. TC1"])
        h = wb["1. TC1"]
        self.assertEqual((h["C6"].value, h["D8"].value, h["H9"].value, h["D10"].value),
                         ("R-1", "Persona de prueba", "prueba@example.com", "X"))
        self.assertEqual((h["E12"].value, h["D13"].value, h["D14"].value, h["H14"].value),
                         ("EMPRESA DE PRUEBA SA", "A00000000", "ES0000000000000000XX", "TC1"))
        self.assertEqual([h[c + "21"].value for c in "DEFGHI"], [100, 100, 100, 100, 100, 200])
        self.assertEqual((h["B25"].value, h["C25"].value, h["H25"].value), ("Ene / Jan", 100, 50))
        self.assertEqual(h["I25"].value, "=SUM(C25:H25)")
        self.assertIn("B23:I23", [str(m) for m in h.merged_cells.ranges])   # sin DH al lado
        self.assertEqual(len(h._charts), 1)
        r = wb["Resumen"]
        self.assertEqual([r.cell(4, n).value for n in range(1, 9)],
                         ["GRUPO DE PRUEBA", "TC1", "ES0000000000000000XX", "EMPRESA DE PRUEBA SA",
                          "Calle Falsa 1", "A00000000", "6.1TD", "31/12/2026"])
        self.assertEqual((r["A2"].value, r["I2"].value, r["O2"].value),
                         ("Cliente/Grupo", "POTENCIAS A CONTRATAR (kW)", "PREVISIÓN DE CONSUMO ANUAL"))
        self.assertEqual([r.cell(4, n).value for n in range(9, 15)], [100] * 5 + [200])
        self.assertEqual(r["O4"].value, "='1. TC1'!C37")
        self.assertEqual(r["U4"].value, "='1. TC1'!I37")
        # fila TOTAL bajo el ultimo suministro
        self.assertEqual((r["T5"].value, r["U5"].value), ("TOTAL", "=SUM(U4:U4)"))
        self.assertEqual([r.cell(5, n).value for n in range(1, 20)], [None] * 19)
        self.assertFalse([m for m in r.merged_cells.ranges if m.min_row == 5])

    def test_consideraciones_en_la_hoja(self):
        c = {"indexado": ["Indexada Omie Pass Through", "Indexada Omie Pass Pool"],
             "inicio": "01/01/2027", "duracion": "31/12/2027", "duracion2": "31/12/2028",
             "temporada": False, "exencion": False, "gdo": True, "excedentes": False}
        cab = {"recepcion_desde": "22/10/2026", "recepcion_hasta": "23/10/2026",
               "recepcion_hora": "12:00"}
        h = openpyxl.load_workbook(io.BytesIO(so.excel_solicitud(
            cab, [{"cups": "ES0000000000000000XX", "tarifa": "2.0TD",
                   "consideraciones": c}])))["1. 000000XX"]
        valores = {celda.value for fila in h.iter_rows(min_row=61) for celda in fila}
        self.assertEqual(h["B61"].value[:32], "Tipología de ofertas solicitadas")
        # el precio fijo sale de la tarifa
        self.assertEqual((h["B62"].value, h["D62"].value), ("Precio fijo / Fixed price",
                                                            "Binómica en DH2"))
        self.assertEqual((h["D63"].value, h["D64"].value, h["D65"].value),
                         ("Indexada Omie Pass Through", "Indexada Omie Pass Pool", None))
        self.assertEqual(h["E67"].value, "Anual / Annual")
        self.assertEqual(h["E70"].value, "01/01/2027")
        self.assertEqual((h["F73"].value, h["F74"].value), ("31/12/2027", "31/12/2028"))
        self.assertEqual((h["B76"].value[:21], h["E76"].value),
                         ("Garantías de origen (", "Sí / Yes"))
        self.assertEqual((h["F79"].value, h["F80"].value), ("22/10/2026", "23/10/2026"))
        self.assertEqual(h["G80"].value, "antes de las 12:00 h / before 12:00")
        self.assertEqual(h["F80"].fill.fgColor.rgb[-6:], so.NARANJA)
        self.assertFalse(any("exención" in str(v) or "excedentes" in str(v) for v in valores))

    def test_sin_precio_fijo_la_fila_queda_vacia(self):
        h = openpyxl.load_workbook(io.BytesIO(so.excel_solicitud({}, [
            {"cups": "ES0000000000000000XX", "tarifa": "6.1TD",
             "consideraciones": {"fijo": False}}])))["1. 000000XX"]
        self.assertEqual((h["B62"].value, h["D62"].value), ("Precio fijo / Fixed price", None))

    def test_precio_fijo_segun_tarifa(self):
        self.assertEqual([so.precio_fijo(t) for t in ("2.0TD", "3.0TD", "6.1TD", "6.4TD")],
                         ["Binómica en DH2"] + ["Binómica en DH6"] * 3)


class TablaConsideraciones(unittest.TestCase):
    def test_comunes_y_cambio_por_suministro(self):
        from ofertas import paginas as op
        comun = {o: False for o in op.COLUMNAS_CASILLA}
        comun.update({"PVPC": True, op.INICIO: dt.date(2027, 1, 1), op.DURACION: None,
                      op.DURACION2: None, op.DURACION3: None, op.PCT_CONSUMO: None, op.PCT_TIPO: None})
        t = op.tabla_consideraciones([{"cups": "A"}, {"cups": "B"}], comun)
        t.loc[1, op.GDO] = True
        t.loc[1, op.EXENCION] = True
        t.loc[1, op.PCT_CONSUMO] = 50.0
        a, b = (op.consideraciones_de(f) for _, f in t.iterrows())
        self.assertEqual((a["indexado"], a["inicio"], a["gdo"]), (["PVPC"], "01/01/2027", False))
        self.assertEqual((a["exencion"], a["pct_consumo_exento"]), (False, None))
        self.assertEqual((b["gdo"], b["exencion"], b["pct_consumo_exento"], b["duracion"]),
                         (True, True, 50.0, ""))


class GuardarYCargar(unittest.TestCase):
    def test_formulario_ida_y_vuelta(self):
        from ofertas import paginas as op
        for v in (dt.date(2027, 1, 1), dt.time(12, 0), "R-1", True, 85.0, None):
            self.assertEqual(op.decodificar(json.loads(json.dumps(op.codificar(v)))), v)
        c = {"fijo": False, "indexado": ["PVPC"], "inicio": "01/01/2027", "duracion": "",
             "duracion2": "31/12/2028", "duracion3": "", "temporada": True, "gdo": True,
             "excedentes": False, "exencion": True, "pct_consumo_exento": 85.0,
             "pct_tipo_exencion": None}
        fila = pd.Series(op.fila_de(c))
        fila[op.INICIO] = pd.Timestamp(fila[op.INICIO])
        fila[op.DURACION2] = pd.Timestamp(fila[op.DURACION2])
        self.assertEqual(op.consideraciones_de(fila), c)

    def test_tabla_con_lo_guardado_si_lo_comun_no_cambia(self):
        from ofertas import paginas as op
        comun = {o: False for o in op.COLUMNAS_CASILLA}
        comun.update({op.INICIO: None, op.DURACION: dt.date(2027, 12, 31), op.DURACION2: None,
                      op.DURACION3: None, op.PCT_CONSUMO: None, op.PCT_TIPO: None})
        fichas = [{"cups": "A", "consideraciones": {"indexado": ["PVPC"],
                                                    "duracion": "31/12/2029"}}]
        guardada = op.tabla_consideraciones(fichas, comun, dict(comun))
        self.assertTrue(guardada.loc[0, "PVPC"])
        self.assertEqual(guardada.loc[0, op.DURACION].year, 2029)
        otra = dict(comun, **{op.GDO: True})               # lo comun ha cambiado
        self.assertFalse(op.tabla_consideraciones(fichas, otra, dict(comun)).loc[0, "PVPC"])


class Correo(unittest.TestCase):
    def suministro(self, cups, tarifa="6.1TD", **c):
        base = {"indexado": ["PVPC"], "inicio": "01/01/2027", "duracion": "31/12/2027",
                "duracion2": "", "duracion3": "", "temporada": False, "exencion": False,
                "gdo": False, "excedentes": False}
        base.update(c)
        return {"grupo": "GRUPO DE PRUEBA", "razon": "EMPRESA DE PRUEBA SA", "cups": cups,
                "tarifa": tarifa, "perfil": {1: [1000, 0, 0, 0, 0, 500]},
                "consideraciones": base}

    cab = {"email": "prueba@example.com", "recepcion_desde": "22/10/2026",
           "recepcion_hasta": "23/10/2026", "recepcion_hora": "12:00"}

    def test_igual_en_todos(self):
        texto = correo.texto_correo(self.cab, [
            self.suministro("ES0000000000000000XX", gdo=True, duracion2="31/12/2028",
                            duracion3="31/12/2029"),
            self.suministro("ES0000000000000001XX", gdo=True, duracion2="31/12/2028",
                            duracion3="31/12/2029")])
        self.assertIn("del cliente GRUPO DE PRUEBA en España", texto)
        self.assertIn("• CUPS: ES0000000000000000XX (EMPRESA DE PRUEBA SA)\n", texto)
        self.assertIn("consumo anual previsto total de 3.000 kWh", texto)
        self.assertIn("• Tipología de precios: Precio fijo y PVPC.\n", texto)
        self.assertIn("• Tipo de contrato: anual.\n", texto)
        self.assertIn("inicio del contrato: 1 de enero de 2027.\n", texto)
        self.assertIn("Duración del contrato: tres alternativas, hasta el 31 de diciembre de "
                      "2027, hasta el 31 de diciembre de 2028 y hasta el 31 de diciembre de "
                      "2029.", texto)
        self.assertIn("deben incluir garantías de origen (GdO).", texto)
        self.assertIn("Fecha máxima: 23 de octubre de 2026, antes de las 12:00 horas.", texto)
        self.assertIn("\nprueba@example.com\n", texto)
        self.assertNotIn("inómica", texto)              # el precio fijo no se menciona
        self.assertNotIn("—", texto)

    def test_exencion_y_excedentes_junto_al_cups(self):
        texto = correo.texto_correo(self.cab, [
            self.suministro("ES0000000000000000XX", exencion=True, pct_consumo_exento=85.0,
                            excedentes=True),
            self.suministro("ES0000000000000001XX", exencion=True, pct_consumo_exento=85.0)])
        self.assertIn("ES0000000000000000XX (EMPRESA DE PRUEBA SA) — sujeto a exención del "
                      "impuesto eléctrico (85 % del consumo); con compensación de excedentes\n",
                      texto)
        self.assertIn("ES0000000000000001XX (EMPRESA DE PRUEBA SA) — sujeto a exención del "
                      "impuesto eléctrico (85 % del consumo)\n", texto)
        resumen = texto.split("Resumen de la solicitud:")[1].split("Adjuntamos")[0]
        self.assertNotIn("exención", resumen)
        self.assertNotIn("excedentes", resumen)

    def test_lo_distinto_en_el_resumen_con_sus_cups(self):
        pt = ["Indexada Omie Pass Through", "Indexada Omie Pass Pool"]
        texto = correo.texto_correo(self.cab, [
            self.suministro("ES0000000000000000XX", indexado=pt, gdo=True),
            self.suministro("ES0000000000000001XX", indexado=pt, gdo=True,
                            inicio="01/03/2027"),
            self.suministro("ES0000000000000002XX", indexado=["PVPC"], temporada=True,
                            fijo=False)])
        self.assertIn("• Tipología de precios: Precio fijo, Indexada Omie Pass Through y Indexada "
                      "Omie Pass Pool (ES0000000000000002XX: PVPC).", texto)
        self.assertIn("• Tipo de contrato: anual (ES0000000000000002XX: de temporada).", texto)
        self.assertIn("inicio del contrato: 1 de enero de 2027 (ES0000000000000001XX: 1 de marzo "
                      "de 2027).", texto)
        self.assertIn("garantías de origen (GdO), salvo en ES0000000000000002XX.", texto)
        self.assertNotIn("—", texto)


def factura_gas(anio, mes, kwh, qd=None):
    fin = (dt.date(anio + mes // 12, mes % 12 + 1, 1) - dt.timedelta(days=1))
    return {"Inicio factura": dt.date(anio, mes, 1), "Fin factura": fin, "Total kWh": kwh,
            "_qd": qd}


def suministro_gas(cups="ES0000000000000000XX", kwh=100_000, tarifa="RL.5", **tecnicos):
    perfil = so.perfil_gas([factura_gas(2026, m, kwh, qd=500) for m in range(1, 13)])
    t = so.tecnicos_por_defecto(sum(p["kwh"] for p in perfil.values()), 500)
    t.update(tecnicos)
    return {"grupo": "GRUPO DE PRUEBA", "alias": "G1", "cups": cups,
            "razon": "EMPRESA DE PRUEBA SA", "direccion": "Calle Falsa 1", "cif": "A00000000",
            "tarifa": tarifa, "fecha_inicio": "01/01/2027", "tipo": so.GAS,
            "perfil_gas": perfil, "tecnicos": t,
            "consideraciones": {"fijo": True, "indexado": ["Indexado TTF con coberturas"],
                                "inicio": "01/01/2027", "duracion": "31/12/2027",
                                "gdo": False, "ieh": so.IEH_REDUCIDO, "pct_ieh": None}}


CAB_GAS = {"ref": "R-G", "producto": so.PRODUCTOS[so.GAS], "producto_tipo": so.GAS,
           "email": "prueba@example.com", "recepcion_desde": "22/10/2026",
           "recepcion_hasta": "23/10/2026", "recepcion_hora": "12:00"}


class GasDatos(unittest.TestCase):
    def test_tarifas_rl(self):
        self.assertEqual([so.grupo_rl(t) for t in ("RL.6", "RL06", "RLTB.6", "RL.11", "6.1TD")],
                         [6, 6, 6, 11, None])
        self.assertEqual((so.tarifa_gas("RL.6"), so.tarifa_gas("RL.10")), ("RL06", "RL10"))
        self.assertEqual([so.grupo_rl_de_consumo(c) for c in (5_000, 5_001, 4_200_000,
                                                              5_000_001, 6e8)], [1, 2, 6, 7, 11])
        self.assertEqual((so.producto_de("gas"), so.producto_de("electricitat"),
                          so.producto_de(None)), (so.GAS, so.ELECTRICIDAD, so.ELECTRICIDAD))

    def test_perfil_gas(self):
        perfil = so.perfil_gas([factura_gas(2025, 12, 31_000, qd=1_200),
                                factura_gas(2026, 1, 62_000, qd=1_300)])
        self.assertEqual(perfil[12], {"anio": 2025, "kwh": 31_000.0, "dias": 31, "qd": 1_200.0})
        self.assertEqual((perfil[1]["anio"], perfil[1]["qd"]), (2026, 1_300.0))
        self.assertEqual(perfil[6], {"anio": None, "kwh": 0.0, "dias": 0, "qd": None})
        self.assertEqual(so.consumo_anual({"perfil_gas": perfil}), 93_000.0)

    def test_telemedida_y_termino_fijo(self):
        # hasta 5 GWh: sin obligacion de telemedida, cuota fija (por cliente)
        t = so.tecnicos_por_defecto(4_200_000, 12_000)
        self.assertEqual((t["telemedida"], t["termino_fijo"], t["qd"]),
                         (False, so.CUOTA_FIJA, 12_000))
        self.assertNotIn("alquiler", t)
        self.assertEqual(so.reglas_tecnicas(t, 4_200_000), (t, []))
        # mas de 5 GWh: telemedida obligatoria y termino variable aunque se cambie a mano
        t2, cambios = so.reglas_tecnicas(dict(t, telemedida=False), 6_000_000)
        self.assertEqual((t2["telemedida"], t2["termino_fijo"], len(cambios)),
                         (True, so.VARIABLE, 2))
        self.assertTrue(so.tecnicos_por_defecto(6_000_000)["telemedida"])

    def test_cogeneracion_sin_tipo_reducido(self):
        self.assertEqual(so.ieh_valido(so.IEH_REDUCIDO, True), so.IEH_EXENTO)
        self.assertEqual(so.ieh_valido(so.IEH_GENERAL, True), so.IEH_GENERAL)
        self.assertEqual(so.ieh_valido(so.IEH_REDUCIDO, False), so.IEH_REDUCIDO)

    def test_avisos(self):
        self.assertEqual(so.avisos_gas(suministro_gas()), [])        # 1,2 GWh en RL.5
        s = suministro_gas(tarifa="RL.6")
        self.assertIn("no corresponde", so.avisos_gas(s)[0])
        s = suministro_gas(telemedida=True, termino_fijo=so.VARIABLE, qd=5_000)
        self.assertEqual(len(so.avisos_gas(s)), 0)
        s["tecnicos"]["qd"] = 4_000            # 1,2 GWh / (4.000 x 365) = 82 %
        self.assertIn("factor de carga del 82 %", so.avisos_gas(s)[0])
        s = suministro_gas(termino_fijo=so.VARIABLE)
        self.assertIn("variable sin telemedida", so.avisos_gas(s)[0])
        s = suministro_gas()
        s["consideraciones"]["duracion"] = "30/06/2027"
        self.assertIn("menos de 12 meses", so.avisos_gas(s)[0])
        s = suministro_gas(interrumpible=True)
        self.assertEqual(so.avisos_gas(s), ["interrumpible sin cumplir los requisitos: más de "
                                            "10 GWh/año, más de 26.000 kWh/día, más de 4 bar, "
                                            "telemedida"])

    def test_facturas_de_gas_de_gemweb(self):
        import xml.etree.ElementTree as ET
        from gemweb_extractor import paginas as gw
        raiz = ET.fromstring(
            "<r><factura><periode_ini>01/09/2026</periode_ini><periode_fin>30/09/2026"
            "</periode_fin><tarifa>RL06</tarifa><consum_tot>300000</consum_tot><consum_p1/>"
            "<qd_aplicada>12000</qd_aplicada><import_lloguer>150.00</import_lloguer>"
            "</factura></r>")
        filas, descartadas = gw.facturas_por_periodo("ES0000000000000000XX", raiz)
        self.assertEqual(descartadas, 0)
        f = filas[0]
        self.assertEqual((f["Total kWh"], f["_qd"], f["P1 kWh"]), (300_000.0, 12_000.0, 0.0))


class GasExcel(unittest.TestCase):
    def test_hoja_y_resumen(self):
        s = suministro_gas(presion=None, cogeneracion=True)
        s["consideraciones"].update(ieh=so.IEH_EXENTO, pct_ieh=60.0, gdo=True)
        wb = openpyxl.load_workbook(io.BytesIO(so.excel_solicitud(CAB_GAS, [s])))
        h = wb["1. G1"]
        self.assertEqual(h["D10"].value, "GAS NATURAL / NATURAL GAS")
        self.assertEqual(h["B18"].value, "Aspectos técnicos / Technical aspects")
        self.assertEqual((h["D19"].value, h["G19"].value, h["I19"].value),
                         ("RL05", "-", "=E40"))
        self.assertEqual((h["D20"].value, h["D21"].value, h["I21"].value),
                         ("SÍ", "NO", "-"))
        self.assertEqual((h["D22"].value, h["E22"].value, h["D23"].value, h["I23"].value),
                         (500, None, "NO", "NO"))
        self.assertEqual(h["D24"].value, "Cuota fija / Fixed fee")
        self.assertEqual(h["D24"].fill.fgColor.rgb[-6:], so.AMARILLO)
        # perfil: mes con su año, Qd media (kWh/dias), energia y Qd aplicada; total en E40
        self.assertEqual((h["B28"].value, h["C28"].value, h["E28"].value, h["G28"].value),
                         ("ene-26", round(100_000 / 31), 100_000, 500))
        self.assertEqual(h["E40"].value, "=SUM(E28:E39)")
        self.assertEqual(len(h._charts), 1)
        valores = [c.value for fila in h.iter_rows(min_row=62) for c in fila if c.value]
        self.assertIn("Indexado TTF (con coberturas / with hedging)", valores)
        self.assertIn("Sí / Yes", valores)                          # precio fijo, sin binomica
        self.assertIn("Exento / Exempt", valores)
        self.assertIn("60 %", valores)
        self.assertTrue(any(str(v).startswith("Garantías de origen de gas renovable")
                            for v in valores))
        self.assertFalse(any("inómica" in str(v) or "eléctrico" in str(v) for v in valores))
        r = wb["Resumen"]
        self.assertEqual([r.cell(4, n).value for n in (7, 9, 10, 11, 12)],
                         ["RL05", 500, "NO", so.CUOTA_FIJA, "='1. G1'!E40"])
        self.assertEqual((r["I2"].value, r["L2"].value), ("ASPECTOS TÉCNICOS",
                                                          "Consumo anual (kWh)"))
        self.assertEqual((r["K5"].value, r["L5"].value), ("TOTAL", "=SUM(L4:L4)"))

    def test_impuesto_reducido_sin_porcentaje(self):
        h = openpyxl.load_workbook(io.BytesIO(so.excel_solicitud(CAB_GAS, [suministro_gas()])))[
            "1. G1"]
        valores = [c.value for fila in h.iter_rows(min_row=62) for c in fila if c.value]
        self.assertIn(f"{so.IEH_REDUCIDO} / Reduced rate (professional use)", valores)
        self.assertFalse(any("renovable" in str(v) for v in valores))


class GasFormulario(unittest.TestCase):
    def test_consideraciones_ida_y_vuelta(self):
        from ofertas import paginas as op
        c = {"fijo": False, "indexado": ["Indexado MIBGAS (PVB)", "Indexado Brent con coberturas"],
             "inicio": "01/01/2027", "duracion": "31/12/2027", "duracion2": "", "duracion3": "",
             "temporada": False, "gdo": True, "ieh": so.IEH_GENERAL, "pct_ieh": 40.0}
        fila = pd.Series(op.fila_de_gas(c))
        for col in (op.INICIO, op.DURACION):
            fila[col] = pd.Timestamp(fila[col])
        self.assertEqual(op.consideraciones_de_gas(fila), c)

    def test_tabla_de_aspectos_tecnicos(self):
        from ofertas import paginas as op
        s = suministro_gas(presion=6.0, dias=5)
        t = op.tabla_tecnicos([s])
        self.assertEqual((t.loc[0, "Tarifa"], t.loc[0, "Consumo anual kWh"]),
                         ("RL05", 1_200_000))
        self.assertEqual(op.tecnicos_de(t.iloc[0]), s["tecnicos"])
        t.loc[0, op.COLUMNAS_TECNICOS["presion"]] = None
        self.assertIsNone(op.tecnicos_de(t.iloc[0])["presion"])


class GasCorreo(unittest.TestCase):
    def test_correo_de_gas(self):
        a = suministro_gas()
        b = suministro_gas("ES0000000000000001XX", cogeneracion=True, telemedida=True,
                           termino_fijo=so.VARIABLE)
        b["consideraciones"].update(ieh=so.IEH_EXENTO, pct_ieh=60.0)
        texto = correo.texto_correo(CAB_GAS, [a, b])
        self.assertIn("suministros de gas natural del cliente GRUPO DE PRUEBA", texto)
        self.assertIn("ofertas de gas natural para", texto)
        self.assertIn("• CUPS: ES0000000000000000XX (EMPRESA DE PRUEBA SA)\n", texto)
        self.assertIn("ES0000000000000001XX (EMPRESA DE PRUEBA SA) — Impuesto sobre "
                      "Hidrocarburos: exento (60 % del consumo); con cogeneración\n", texto)
        self.assertIn("total de 2.400.000 kWh y un caudal diario (Qd) total de 1.000 kWh/día.",
                      texto)
        self.assertIn("• Tipología de precios: Precio fijo y Indexado TTF con coberturas.", texto)
        self.assertIn("• Impuesto sobre Hidrocarburos: tipo reducido 0,15 €/GJ (usos "
                      "profesionales) (ES0000000000000001XX: exento (60 % del consumo)).", texto)
        self.assertIn("• Término fijo: cuota fija (por cliente) (ES0000000000000001XX: variable "
                      "(por caudal)).", texto)
        self.assertNotIn("electricidad", texto)


if __name__ == "__main__":
    unittest.main()
