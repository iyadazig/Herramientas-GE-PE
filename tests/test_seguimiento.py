# -*- coding: utf-8 -*-
"""Seguimiento y Control: importacion del Excel, edicion del cuadro, tablas auxiliares,
exportacion y calculo del retraso. Todo con un Excel ficticio generado aqui."""

import datetime as dt
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import almacen
import config
from seguimiento import columnas as C
from seguimiento import datos as D
from seguimiento import retraso as R

CUPS1, CUPS2, CUPS3 = ("ES0000000000000000%sXX" % n for n in "123")


def excel_ficticio():
    """Excel con la estructura del cuadro real y datos inventados."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CLIENTES"
    meses = []
    for m in range(1, 13):
        meses += [dt.datetime(2026, m, 1), "Obs"]
    ws.append(["Técnico", "Auxiliar", "Otra información", "Tipo Contratación",
               "Cadencia informe", "Cliente/Grupo", "CIF", "Titular de contrato", "CUPS",
               "Denominación suministro", "Tipo suministro (Electricidad / Gas)",
               "Tarifa de acceso", "Max. potencia P3/P6 (kW)", "Consumo anual (kWh)",
               "Fecha denuncia contrato", "Fecha fin contrato",
               "Fecha prevista lanzamiento licitación", "Comercializadora",
               "Tipo de contrato (Precio fijo / Indexado / Indexado+coberturas)", "Provincia",
               "Gestor Comercial", "Servicio de SyC (€/mes)", "Servicio SyC FV (€/mes)",
               "Servicio SGE (€/mes)", "Servicio Seg. Mercados (€/mes)",
               "Total cuota SyC (€/mes)", "Servicio SMAC (€/mes)", "RENOVACIÓN ABIERTA"]
              + meses + [None, "Observaciones", "Observaciones Extra", "Validación Gemweb",
                         "BAJA EN GEYPE"])
    estados1 = ["Inf enviado", "-", "Inf enviado", None, "Recibida", None] + [None] * 18
    ws.append(["Tecnico Uno", "Aux Uno", None, "Sgto. y Control", "Mensual", "Grupo Uno",
               "B00000000", "Titular Uno", CUPS1, "Planta Uno", "Electricidad", "3.0TD", 50,
               "1.425.738\xa0", dt.datetime(2027, 1, 1), dt.datetime(2027, 1, 31), None,
               "Comercializadora X", "Indexado  ", "Sevilla", "Gestor A", 10.5, 2, None, None,
               12.5, None, "NO"] + estados1 + [None, "Clave", None, "SI", "NO"])
    estados2 = []
    for m in range(1, 13):
        estados2 += ["Pdte inf" if m <= 6 else "Sin fra", None]
    ws.append(["Tecnico Dos", "Aux Uno", None, "Sgto. y Control", "Semestral", "Grupo Dos",
               "B00000001", "Titular Dos", CUPS2, "Planta Dos", "Gas", "RL4", "-", 1000,
               "NO PRÓRROGA", " 31/12/2027", None, "Comercializadora Y", "Fijo", "Cádiz",
               "Gestor A", 20, None, None, None, 20, None, "SI"] + estados2
              + [None, None, None, "SI", "NO"])
    ws.append(["Tecnico Uno", "Aux Uno", None, "Asesoramiento", "Mensual", "Grupo Uno",
               "B00000000", "Titular Uno", CUPS3, "Planta Tres", "Electricidad", "2.0TD", 10,
               500, None, None, None, "Comercializadora X", "Fijo", "Sevilla", "Gestor A", 5,
               None, None, None, 5, None, "NO"] + [None] * 24 + [None, None, None, None, None])
    ws.append([None] * 10)                                           # fila vacia
    bajas = wb.create_sheet("BAJAS")
    bajas.append(["Técnico", "Auxiliar", "Cadencia revisión", "Cadencia informe",
                  "Cliente/Grupo", "CUPS", "Denominación suministro", dt.datetime(2026, 1, 1),
                  "Obs"])
    bajas.append(["Tecnico Uno", "Aux Uno", "Mensual", "Anual", "Grupo Viejo",
                  "ES0000000000000000BAJA", "Planta Vieja", "Sin fra", None])
    claves = wb.create_sheet("CLAVES")
    claves.append(["Técnico", "Auxiliar", "Cliente/Grupo", "Denominación suministro",
                   "Usuario", "Clave", "Comercializadora", "Enlace web", "Gestor", "email",
                   "Teléfono", None])
    claves.append(["Tecnico Uno", "Aux Uno", "Grupo Uno", "Planta Uno", "usuario1",
                   "clave-ficticia", "Comercializadora X", "https://ejemplo.invalid", None,
                   None, 600000000, "nota suelta"])
    contactos = wb.create_sheet("CONTACTOS")
    contactos.append(["Técnico", "Auxiliar", "Cliente/Grupo", "Denominación suministro",
                      "Contacto", "Puesto", "email", "Teléfono", "Comentarios"])
    contactos.append(["Tecnico Uno", "Aux Uno", "Grupo Uno", "Planta Uno", "Persona Uno",
                      "Gerente", "persona@ejemplo.invalid", 955000000, None])
    smac = wb.create_sheet("SMAC")
    smac.append(["Técnico", "Auxiliar", "Código Pedido", "Cliente/Grupo",
                 "Denominación suministro", "Usuario", "Clave", "Último dato", "Observaciones"])
    tele = wb.create_sheet("TELEMEDIDA")
    tele.append(["Técnico", "Auxiliar", "Cliente/Grupo", "Denominación suministro", "TELÉFONO",
                 "D. ENLACE", "P.M.", "C.L.", "MODEM", "IP", "PARIDAD", "VELOCIDAD", "TÉCNICO",
                 "OBSERVACIONES"])
    tele.append(["Tecnico Dos", "Aux Uno", None, "Planta Dos", 600000001, 1, 1, 1, None, None,
                 None, None, "Instalador", None])
    lista = wb.create_sheet("Lista")
    for v, d, c in C.ESTADOS:
        lista.append([None, v, d, c])
    lista.append(["TÉCNICOS"])
    lista.append(["Tecnico Uno"])
    lista.append(["AUXILIARES"])
    lista.append(["Aux Uno"])
    lista.append(["TIPOS DE CONTRATO"])
    lista.append(["Fijo"])
    lista.append(["Indexado"])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def contactos_ficticios():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Técnico", "Auxiliar", "Cliente/Grupo", "Denominación suministro", "CUPS",
               "Fecha de baja", "Nombre cliente", "Para", "CC"])
    ws.append(["Tecnico Uno", "Aux Uno", "Grupo Uno", "Planta Uno", CUPS1, None, "Cliente Uno",
               "uno@ejemplo.invalid", "copia@ejemplo.invalid"])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


class BaseDatos(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.p = mock.patch.object(config, "CARPETA_DATOS", Path(self.tmp.name))
        self.p.start()
        almacen._iniciado.clear()

    def tearDown(self):
        self.p.stop()
        almacen._iniciado.clear()
        self.tmp.cleanup()

    def importar(self):
        leido = D.leer_cuadro(excel_ficticio())
        D.importar(leido, D.leer_contactos_envio(contactos_ficticios()), "prueba")
        return leido


class Valores(unittest.TestCase):
    def test_numeros_con_formato_espanol(self):
        self.assertEqual(D.numero("1.425.738\xa0"), 1425738)
        self.assertEqual(D.numero("23,393"), 23.393)
        self.assertEqual(D.numero("26.82"), 26.82)
        self.assertEqual(D.numero("695.731"), 695731)
        self.assertIsNone(D.numero("-"))

    def test_fechas_escritas_como_texto(self):
        self.assertEqual(D.fecha(" 30/09/2027"), "2027-09-30")
        self.assertEqual(D.fecha("30/06/02027"), "2027-06-30")
        self.assertIsNone(D.fecha("NO PRÓRROGA"))

    def test_estados_en_su_forma_canonica(self):
        self.assertEqual(D.estado_canonico("inf ENVIADO "), C.INF_ENVIADO)
        self.assertEqual(D.estado_canonico("-"), "-")


class Importacion(BaseDatos):
    def test_lee_todas_las_pestanas(self):
        leido = self.importar()
        self.assertEqual(leido["recuento"]["CLIENTES"], 3)
        self.assertEqual(leido["recuento"]["BAJAS"], 1)
        activos = almacen.listar_sgc_suministros(bajas=False)
        self.assertEqual(len(activos), 3)
        self.assertEqual(len(almacen.listar_sgc_suministros(bajas=True)), 1)
        self.assertEqual(len(almacen.listar_sgc_filas("claves")), 1)
        self.assertEqual(len(almacen.listar_sgc_filas("contactos_envio")), 1)
        uno = next(s for s in activos if s["cups"] == CUPS1)["datos"]
        self.assertEqual(uno["Consumo anual (kWh)"], 1425738)
        self.assertEqual(uno["Tipo de contrato"], "Indexado")          # espacios quitados
        self.assertEqual(uno["Fecha fin contrato"], "2027-01-31")

    def test_lo_que_no_encaja_va_a_otros(self):
        self.importar()
        dos = next(s for s in almacen.listar_sgc_suministros() if s["cups"] == CUPS2)["datos"]
        self.assertIn("Fecha denuncia contrato: NO PRÓRROGA", dos[C.OTROS])
        self.assertIsNone(dos["Max. potencia P3/P6 (kW)"])                # «-» es vacio
        baja = almacen.listar_sgc_suministros(bajas=True)[0]["datos"]
        self.assertIn("Cadencia revisión: Mensual", baja[C.OTROS])
        clave = almacen.listar_sgc_filas("claves")[0]["datos"]
        self.assertEqual(clave["Teléfono"], "600000000")
        self.assertEqual(clave[C.OTROS], "nota suelta")

    def test_telemedida_distingue_tecnico_de_su_tecnico(self):
        self.importar()
        t = almacen.listar_sgc_filas("telemedida")[0]["datos"]
        self.assertEqual(t["Técnico"], "Tecnico Dos")
        self.assertEqual(t["Técnico telemedida"], "Instalador")

    def test_estados_por_mes_y_listas(self):
        self.importar()
        estados = almacen.listar_sgc_estados(2026)
        uno = next(s["id"] for s in almacen.listar_sgc_suministros() if s["cups"] == CUPS1)
        self.assertEqual(estados[(uno, 2026, 1)]["estado"], C.INF_ENVIADO)
        self.assertEqual(estados[(uno, 2026, 1)]["obs"], "-")
        listas = almacen.listar_sgc_listas()
        self.assertEqual([v["valor"] for v in listas["tipos_contrato"]], ["Fijo", "Indexado"])
        self.assertEqual(listas["estados"][0]["color"], "#FFFFFF")
        equipo = {e["nombre"]: e for e in almacen.listar_sgc_equipo()}
        self.assertTrue(equipo["Tecnico Dos"]["tecnico"])
        self.assertTrue(equipo["Aux Uno"]["auxiliar"])

    def test_reimportar_conserva_emails_del_equipo_y_contactos(self):
        self.importar()
        equipo = almacen.listar_sgc_equipo()
        for e in equipo:
            e["email"] = "equipo@ejemplo.invalid" if e["nombre"] == "Aux Uno" else None
        almacen.guardar_sgc_equipo(equipo)
        D.importar(D.leer_cuadro(excel_ficticio()), None, "prueba")
        aux = next(e for e in almacen.listar_sgc_equipo() if e["nombre"] == "Aux Uno")
        self.assertEqual(aux["email"], "equipo@ejemplo.invalid")
        self.assertEqual(len(almacen.listar_sgc_filas("contactos_envio")), 1)


class Edicion(BaseDatos):
    def test_editar_anadir_y_borrar_en_el_cuadro(self):
        self.importar()
        original = D.cuadro_df(2026)
        editado = original.copy()
        i = editado.index[editado["CUPS"] == CUPS1][0]
        editado.at[i, "Observaciones"] = "Cambiada"
        editado.at[i, "feb-26"] = C.PDTE_INF
        editado = editado[editado["CUPS"] != CUPS3]                     # borrar -> baja
        nueva = {k: None for k in editado.columns}
        nueva.update({"CUPS": "ES0000000000000000NUEVA", "Técnico": "Tecnico Uno",
                      "mar-26": C.RECIBIDA})
        editado = pd.concat([editado, pd.DataFrame([nueva])], ignore_index=True)
        n, e, b = D.guardar_cuadro(original, editado, 2026, "prueba")
        self.assertEqual((n, e, b), (2, 2, 1))
        activos = {s["cups"]: s for s in almacen.listar_sgc_suministros(bajas=False)}
        self.assertEqual(activos[CUPS1]["datos"]["Observaciones"], "Cambiada")
        self.assertEqual(activos[CUPS1]["datos"]["Consumo anual (kWh)"], 1425738)  # intacto
        self.assertNotIn(CUPS3, activos)
        estados = almacen.listar_sgc_estados(2026)
        self.assertEqual(estados[(activos[CUPS1]["id"], 2026, 2)]["estado"], C.PDTE_INF)
        nuevo = activos["ES0000000000000000NUEVA"]["id"]
        self.assertEqual(estados[(nuevo, 2026, 3)]["estado"], C.RECIBIDA)
        baja = next(s for s in almacen.listar_sgc_suministros(bajas=True) if s["cups"] == CUPS3)
        self.assertEqual(baja["datos"]["Fecha de baja"], dt.date.today().isoformat())
        D.dar_de_baja([baja["id"]], "prueba", baja=False)
        self.assertIn(CUPS3, {s["cups"] for s in almacen.listar_sgc_suministros(bajas=False)})

    def test_sin_cambios_no_guarda_nada(self):
        self.importar()
        original = D.cuadro_df(2026)
        self.assertEqual(D.guardar_cuadro(original, original.copy(), 2026, "prueba"), (0, 0, 0))

    def test_columnas_ocultas_no_se_borran(self):
        self.importar()
        original = D.cuadro_df(2026)[[D.ID, "CUPS", "Observaciones"]]
        editado = original.copy()
        editado.loc[editado["CUPS"] == CUPS2, "Observaciones"] = "Nueva"
        D.guardar_cuadro(original, editado, 2026, "prueba")
        dos = next(s for s in almacen.listar_sgc_suministros() if s["cups"] == CUPS2)
        self.assertEqual(dos["datos"]["Tarifa de acceso"], "RL4")
        self.assertEqual(almacen.listar_sgc_estados(2026)[(dos["id"], 2026, 1)]["estado"],
                         C.PDTE_INF)

    def test_tabla_auxiliar(self):
        self.importar()
        original = D.tabla_df("claves")
        editado = original.copy()
        editado.loc[0, "Clave"] = "otra-ficticia"
        editado = pd.concat([editado, pd.DataFrame([{"Cliente/Grupo": "Grupo Dos",
                                                     "Usuario": "u2"}])], ignore_index=True)
        self.assertEqual(D.guardar_tabla("claves", original, editado, "prueba"), (2, 0))
        claves = [f["datos"] for f in almacen.listar_sgc_filas("claves")]
        self.assertEqual(claves[0]["Clave"], "otra-ficticia")
        self.assertEqual(claves[0][C.OTROS], "nota suelta")              # columna no editada
        self.assertEqual(D.guardar_tabla("claves", D.tabla_df("claves"),
                                         D.tabla_df("claves").iloc[1:], "prueba"), (0, 1))

    def test_exportar_y_volver_a_leer(self):
        self.importar()
        contenido = D.excel_completo(2026)
        wb = openpyxl.load_workbook(io.BytesIO(contenido))
        self.assertIn("CLIENTES", wb.sheetnames)
        self.assertIn("CONTACTOS ENVÍO", wb.sheetnames)
        releido = D.leer_cuadro(contenido)
        self.assertEqual(releido["recuento"]["CLIENTES"], 3)
        self.assertEqual(len(releido["filas"]["claves"]), 1)


def _sum(id_, tipo="Sgto. y Control", cadencia="Mensual", tecnico="T", auxiliar="A"):
    return {"id": id_, "datos": {"Tipo contratación": tipo, "Cadencia informe": cadencia,
                                 "Técnico": tecnico, "Auxiliar": auxiliar, "CUPS": "C%d" % id_,
                                 "Cliente/Grupo": "G"}}


def _est(id_, por_mes, anio=2026):
    return {(id_, anio, m): {"estado": e} for m, e in por_mes.items()}


class Retraso(unittest.TestCase):
    def test_tamano_de_bloque_como_la_macro(self):
        self.assertEqual(R.tam_bloque("Mensual"), 1)
        self.assertEqual(R.tam_bloque("Trimestral "), 3)
        self.assertEqual(R.tam_bloque("semestral"), 6)
        self.assertEqual(R.tam_bloque("Anual"), 12)
        self.assertEqual(R.tam_bloque(None), 0)

    def test_fecha_de_control_m_mas_2(self):
        self.assertEqual(R.fecha_control(2026, 1), dt.date(2026, 3, 1))
        self.assertEqual(R.fecha_control(2026, 12), dt.date(2027, 2, 1))

    def test_mensual(self):
        estados = _est(1, {1: "Inf enviado", 2: "Recibida", 3: "Sin fra", 4: "Inf no enviado",
                           5: "Temporal", 6: "Cadencia"})
        filas = R.auditoria([_sum(1)], estados, 2026, dt.date(2026, 10, 26))
        self.assertEqual([(f["Bloque"], f["Resultado"]) for f in filas],
                         [(2, R.RETRASO), (4, R.RETRASO), (1, R.ENVIADO)])

    def test_bloque_aun_no_vencido(self):
        estados = _est(1, {8: "Recibida", 9: "Pdte inf"})
        filas = R.auditoria([_sum(1)], estados, 2026, dt.date(2026, 10, 26))
        self.assertEqual([f["Bloque"] for f in filas], [8])              # sep vence el 1/11

    def test_trimestral_y_semestral(self):
        estados = _est(1, {1: "Inf enviado", 2: "Inf enviado", 3: "Sin fra",
                           4: "Inf enviado", 5: "Recibida"})
        estados.update(_est(2, {m: "Inf enviado" for m in range(1, 7)}))
        filas = R.auditoria([_sum(1, cadencia="Trimestral"), _sum(2, cadencia="Semestral")],
                            estados, 2026, dt.date(2026, 10, 26))
        resultado = {(f["_id"], f["Bloque"]): (f["Resultado"], f["Inf. enviados"],
                                               f["Inf. esperados"]) for f in filas}
        self.assertEqual(resultado, {(1, 1): (R.ENVIADO, 2, 2), (1, 2): (R.RETRASO, 1, 2),
                                     (2, 1): (R.ENVIADO, 6, 6)})

    def test_solo_seguimiento_y_control(self):
        estados = _est(1, {1: "Recibida"})
        self.assertEqual(R.auditoria([_sum(1, tipo="Asesoramiento")], estados, 2026,
                                     dt.date(2026, 10, 26)), [])
        self.assertEqual(R.auditoria([_sum(1, cadencia=None)], estados, 2026,
                                     dt.date(2026, 10, 26)), [])

    def test_resumen_e_indicadores(self):
        estados = _est(1, {1: "Inf enviado", 2: "Recibida"})
        estados.update(_est(2, {1: "Inf enviado"}))
        df = R.auditoria_df(R.auditoria(
            [_sum(1, tecnico="T1"), _sum(2, tecnico="", auxiliar="A2")], estados, 2026,
            dt.date(2026, 10, 26)))
        self.assertEqual(R.indicadores(df), (1, 2, 2 / 3))
        tec = R.resumen(df, "Técnico")
        self.assertEqual(tec.to_dict("records"),
                         [{"Técnico": "T1", "Retrasos": 1, "Enviados": 1, "% Retraso": 0.5}])


if __name__ == "__main__":
    unittest.main()
