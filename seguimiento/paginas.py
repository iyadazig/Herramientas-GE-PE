# -*- coding: utf-8 -*-
"""
Paginas del modulo «Seguimiento y Control»
==========================================
Cuadro de seguimiento, retraso de informes, tablas auxiliares (contactos de envio, contactos
y claves de clientes, SMAC, telemedida), bajas y la parte de administracion (importar los
Excel, equipo, listas y Outlook). Los envios de informes y avisos estan en envios.py.
"""

import datetime as dt
import io
import re
import zlib

import altair as alt
import openpyxl
import pandas as pd
import streamlit as st

import almacen
import config
import estilo
from seguimiento import columnas as C
from seguimiento import datos as D
from seguimiento import retraso as R

MODULO = "Seguimiento y Control"
ss = st.session_state
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EMAIL = re.compile(r"[^@\s;,]+@[^@\s;,]+\.[^@\s;,]+")


def _usuario():
    return ss.usuario["usuario"]


def _firma(*partes):
    """Texto corto que identifica un filtro (para la clave del editor: al cambiar el filtro
    el editor empieza de nuevo)."""
    return "%08x" % zlib.crc32(repr(partes).encode("utf-8"))


def _num(v, d=2):
    return ("{:,.%df}" % d).format(v or 0).replace(",", "X").replace(".", ",").replace("X", ".")


def _hay_datos():
    if almacen.listar_sgc_suministros(bajas=False) or almacen.listar_sgc_filas("contactos_envio"):
        return True
    st.info("Todavía no hay datos de Seguimiento y Control. Un administrador tiene que importar "
            "el Excel del cuadro de seguimiento en Administración → «Seguimiento y Control».")
    return False


def _anio():
    anios = sorted(set(almacen.anios_sgc()) | {dt.date.today().year})
    return st.selectbox("Año", anios, index=anios.index(dt.date.today().year), key="sgc_anio")


def _opciones_con_valores(opciones, serie):
    """Opciones de un desplegable mas los valores que ya tiene la columna (para no perderlos)."""
    vistos = list(dict.fromkeys(opciones))
    for v in serie.dropna().unique():
        if v not in vistos:
            vistos.append(v)
    return vistos


def _config_columnas(definiciones, df, listas, equipo, estados=()):
    cfg = {D.ID: None}
    for c in definiciones:
        k = c["clave"]
        if c["tipo"] == C.FECHA:
            cfg[k] = st.column_config.DateColumn(k, format="DD/MM/YYYY")
        elif c["tipo"] == C.NUMERO:
            cfg[k] = st.column_config.NumberColumn(k, format="localized")
        elif c["tipo"] == C.CALCULADA:
            cfg[k] = st.column_config.NumberColumn(k, format="localized", disabled=True,
                                                   help="Suma de las cuotas de SyC")
        elif c.get("lista") and k in df:
            cfg[k] = st.column_config.SelectboxColumn(
                k, options=_opciones_con_valores(D.opciones(listas, equipo, c["lista"]), df[k]))
    for k in estados:
        if k in df:
            cfg[k] = st.column_config.SelectboxColumn(
                k, options=_opciones_con_valores(D.opciones(listas, equipo, "estados"), df[k]))
    return cfg


def _filtrar(df, clave):
    """Filtros comunes: tecnico, auxiliar, cliente/grupo y texto libre."""
    f = st.columns(4)
    tec = f[0].multiselect("Técnico", sorted(df["Técnico"].dropna().unique()), key=clave + "_tec")
    aux = f[1].multiselect("Auxiliar", sorted(df["Auxiliar"].dropna().unique()),
                           key=clave + "_aux")
    cli = f[2].multiselect("Cliente/Grupo", sorted(df["Cliente/Grupo"].dropna().unique()),
                           key=clave + "_cli")
    texto = f[3].text_input("Buscar en todas las columnas", key=clave + "_txt").strip().lower()
    if tec:
        df = df[df["Técnico"].isin(tec)]
    if aux:
        df = df[df["Auxiliar"].isin(aux)]
    if cli:
        df = df[df["Cliente/Grupo"].isin(cli)]
    if texto:
        df = df[df.drop(columns=[D.ID]).astype(str).apply(
            lambda c: c.str.lower().str.contains(texto, regex=False)).any(axis=1)]
    return df, (tuple(tec), tuple(aux), tuple(cli), texto)


def _boton_excel(nombre_fichero, generar, clave):
    """Excel que se prepara al pulsar (generarlo en cada recarga seria lento)."""
    if st.button("Preparar Excel", icon=":material/table_view:", key=clave):
        ss[clave + "_datos"] = generar()
    if ss.get(clave + "_datos"):
        estilo.boton_descarga(MODULO, st, "Descargar %s" % nombre_fichero, ss[clave + "_datos"],
                              nombre_fichero, XLSX)


# ------------------------------------------------------------------- cuadro
def pagina_cuadro():
    if not _hay_datos():
        return
    listas, equipo = almacen.listar_sgc_listas(), almacen.listar_sgc_equipo()
    c1, c2, c3, c4 = st.columns([1, 2, 2, 2])
    with c1:
        anio = _anio()
    df = D.cuadro_df(anio)
    tipo = c2.multiselect("Tipo de contratación", sorted(df["Tipo contratación"].dropna().unique()),
                          key="sgc_tipo")
    cadencia = c3.multiselect("Cadencia", sorted(df["Cadencia informe"].dropna().unique()),
                              key="sgc_cad")
    meses = {C.columna_mes(anio, m): m for m in range(1, 13)}
    mes = c4.selectbox("Estado de un mes", ["—"] + list(meses), key="sgc_mes")
    estado = None
    if mes != "—":
        estado = c4.multiselect("Estados de %s" % mes,
                                D.opciones(listas, equipo, "estados") + ["(vacío)"],
                                key="sgc_estado")
    df, firma = _filtrar(df, "sgc_cuadro")
    if tipo:
        df = df[df["Tipo contratación"].isin(tipo)]
    if cadencia:
        df = df[df["Cadencia informe"].isin(cadencia)]
    if mes != "—" and estado:
        vacio = "(vacío)" in estado
        df = df[df[mes].isin(estado) | (df[mes].isna() if vacio else False)]

    m = st.columns(4)
    m[0].metric("Suministros", len(df))
    m[1].metric("Clientes/Grupos distintos", df["Cliente/Grupo"].nunique())
    m[2].metric("CUPS distintos", df["CUPS"].nunique())
    m[3].metric("Cuota total SyC (€/mes)", _num(df[C.TOTAL_CUOTA].sum(), 0))

    grupos = st.multiselect("Columnas que se muestran", C.GRUPOS_CUADRO,
                            default=[C.CLIENTE, "Meses", C.OBSERVACIONES], key="sgc_grupos")
    de_grupos = {c["clave"] for c in C.CUADRO if c["grupo"] in grupos}
    columnas = [c for c in D.columnas_cuadro(anio) if c in de_grupos or (
        "Meses" in grupos and (c in meses or c.startswith("Obs ")))]
    vista = df[[D.ID] + columnas]
    colores = D.colores_estados(listas)

    pestana = st.radio("Vista", ["Editar", "Ver con colores"], horizontal=True,
                       key="sgc_vista", label_visibility="collapsed")
    if pestana == "Ver con colores":
        def color(v):
            c = colores.get(v)
            if not c:
                return ""
            texto = "#FFFFFF" if c.upper() in ("#FF0000", "#7030A0") else "#1E1E1E"
            return "background-color: %s; color: %s" % (c, texto)
        st.dataframe(vista.drop(columns=[D.ID]).style.map(
            color, subset=[c for c in columnas if c in meses]), hide_index=True,
            use_container_width=True, height=600)
    else:
        st.caption("Edita las celdas y pulsa «Guardar cambios». Para añadir un suministro, usa "
                   "la última fila; si borras una fila, el suministro pasa a «Bajas». Guarda "
                   "antes de cambiar los filtros o el año: al cambiarlos se pierde lo no guardado.")
        cfg = _config_columnas([c for c in C.CUADRO if c["clave"] in columnas], vista, listas,
                               equipo, estados=[c for c in columnas if c in meses])
        editado = st.data_editor(vista, hide_index=True, use_container_width=True,
                                 num_rows="dynamic", column_config=cfg, height=600,
                                 key="sgc_editor_%d_%s" % (anio, _firma(
                                     firma, columnas, tipo, cadencia, mes, estado)))
        if st.button("Guardar cambios", type="primary", icon=":material/save:"):
            n, e, b = D.guardar_cuadro(vista, editado, anio, _usuario())
            almacen.registrar_actividad(_usuario(), MODULO, "Guardar cuadro",
                                        "%d suministros, %d meses, %d bajas" % (n, e, b))
            st.success("Guardado: %d suministros, %d celdas de meses y %d bajas." % (n, e, b))
            st.rerun()

    _gemweb()
    with st.expander("Leyenda de estados"):
        st.dataframe(pd.DataFrame(almacen.listar_sgc_listas().get("estados", [])).rename(
            columns={"valor": "Estado", "descripcion": "Descripción", "color": "Color"}).style.map(
            lambda c: "background-color: %s" % c if c else "", subset=["Color"]),
            hide_index=True)
    _boton_excel("Seguimiento_y_Control_%d.xlsx" % anio, lambda: D.excel_completo(anio),
                 "sgc_excel")


COLUMNAS_CAMBIOS = {"cups": "CUPS", "denominacion": "Denominación suministro", "Mes": "Mes",
                    "antes": "Estado en el cuadro", "despues": "Estado propuesto",
                    "factura": "Factura"}
FECHAS_LISTADO = {c: st.column_config.DateColumn(format="DD/MM/YYYY")
                  for c in ("Emisión", "Periodo inicio", "Periodo fin")}


def _df_cambios(cambios, aplicados=False):
    """Cambios (propuestos o ya aplicados) para verlos y descargarlos."""
    df = pd.DataFrame(cambios, columns=["fecha", "cups", "denominacion", "anio", "mes",
                                        "antes", "despues", "factura"])
    df["Mes"] = [C.columna_mes(int(a), int(m)) for a, m in zip(df["anio"], df["mes"])]
    columnas = dict(COLUMNAS_CAMBIOS)
    if aplicados:
        df["fecha"] = pd.to_datetime(df["fecha"])
        columnas = {"fecha": "Aplicado", **columnas, "antes": "Antes", "despues": "Después"}
    return df[list(columnas)].rename(columns=columnas)


def _df_listado(listado, columnas):
    df = pd.DataFrame(listado, columns=columnas)
    for c in FECHAS_LISTADO:
        df[c] = pd.to_datetime(df[c], errors="coerce").dt.date
    return df


def _gemweb():
    """Facturas de Gemweb: busca los cambios, los muestra para aceptarlos (todos o los
    seleccionados) y descargarlos; aparte, los CUPS del cuadro que no estan en Gemweb."""
    from seguimiento import estados_gemweb as G
    p = G.propuesta()
    pendientes = len(p["cambios"]) if p else 0
    titulo = "Facturas de Gemweb" + (" · %d cambios pendientes de aceptar" % pendientes
                                     if pendientes else "")
    with st.expander(titulo, expanded=bool(pendientes) or bool(ss.get("sgc_gw_abierto"))):
        st.caption("Comprueba que las facturas de Gemweb con periodo de este año estén señaladas "
                   "en el cuadro y propone los cambios; no se aplica nada hasta que los aceptes. "
                   "Solo se recoge un caso: el cuadro dice que no hay factura (mes vacío o «Sin "
                   "fra») y en Gemweb sí la hay, que pasa a «Recibida». Ningún otro estado se "
                   "toca. Solo suministros «Sgto. y Control». También se busca solo a las 7:00 y "
                   "a las 12:00.")
        c1, c2 = st.columns(2)
        if c1.button("Buscar cambios en Gemweb", type="primary", icon=":material/sync:",
                     key="sgc_gw", use_container_width=True):
            with st.spinner("Consultando las facturas de Gemweb…"):
                try:
                    p = G.buscar()
                except Exception as e:                                # noqa: BLE001
                    st.error("No se ha podido consultar Gemweb: %s" % e)
                    return
            G.guardar_propuesta(p)
            almacen.registrar_actividad(_usuario(), MODULO, "Buscar estados en Gemweb",
                                        "%d facturas, %d cambios" % (p["facturas"],
                                                                     len(p["cambios"])))
            ss.sgc_gw_abierto = True
            st.rerun()
        if c2.button("CUPS del cuadro que no están en Gemweb", icon=":material/search_off:",
                     key="sgc_gw_fuera", use_container_width=True):
            with st.spinner("Consultando los suministros de Gemweb…"):
                try:
                    ss.sgc_gw_fuera_datos = G.cups_fuera_de_gemweb()
                except Exception as e:                                # noqa: BLE001
                    st.error("No se ha podido consultar Gemweb: %s" % e)
            ss.sgc_gw_abierto = True
        if ss.get("sgc_gw_fuera_datos") is not None:
            _cups_fuera(ss.sgc_gw_fuera_datos)
        if p:
            _propuesta(G, p)
        aplicados = almacen.listar_cambios_gemweb(500)
        if aplicados:
            st.markdown("**Últimos cambios aplicados desde Gemweb**")
            st.dataframe(_df_cambios(aplicados, aplicados=True), hide_index=True,
                         use_container_width=True, height=220, column_config={
                             "Aplicado": st.column_config.DatetimeColumn(
                                 format="DD/MM/YYYY HH:mm")})


def _cups_fuera(filas):
    df = pd.DataFrame(filas)
    st.markdown("**%d suministros del cuadro cuyo CUPS no está en Gemweb**" % len(df))
    if df.empty:
        return
    st.dataframe(df, hide_index=True, use_container_width=True, height=250)
    c1, c2 = st.columns([1, 1])
    estilo.boton_descarga(MODULO, c1, "Descargar (Excel)",
                          D.excel_tablas([("CUPS fuera de Gemweb", df)]),
                          "CUPS_no_en_Gemweb_%s.xlsx" % dt.date.today().strftime("%Y%m%d"),
                          XLSX)
    if c2.button("Ocultar", key="sgc_gw_fuera_ocultar"):
        ss.sgc_gw_fuera_datos = None
        st.rerun()


def _propuesta(G, p):
    fecha = dt.datetime.fromisoformat(p["fecha"])
    st.markdown("**Última búsqueda: %s**, facturas con periodo desde el %s."
                % (fecha.strftime("%d/%m/%Y a las %H:%M"),
                   dt.date.fromisoformat(p["inicio"]).strftime("%d/%m/%Y")))
    m = st.columns(3)
    m[0].metric("Facturas encontradas", p["facturas"])
    m[1].metric("Descartadas", p["descartadas"],
                help="Anuladas, simuladas, abonos, abonadas y complementarias")
    m[2].metric("Cambios pendientes de aceptar", len(p["cambios"]))
    listado = _df_listado(p["listado"], G.COLUMNAS_LISTADO)
    cambios = _df_cambios(p["cambios"])
    estilo.boton_descarga(
        MODULO, st, "Descargar cambios y facturas encontradas (Excel)",
        D.excel_tablas([("Cambios propuestos", cambios), ("Facturas encontradas", listado)]),
        "Gemweb_cambios_%s.xlsx" % fecha.strftime("%Y%m%d_%H%M"), XLSX,
        key="sgc_gw_excel_%s" % p["fecha"])
    t1, t2 = st.tabs(["Cambios propuestos", "Facturas encontradas"])
    with t1:
        if cambios.empty:
            st.info("Todas las facturas de Gemweb ya están señaladas en el cuadro.")
        else:
            st.caption("Marca los que quieras aplicar, o aplícalos todos.")
            estilo.boton_descarga(
                MODULO, st, "Descargar los %d cambios propuestos (Excel)" % len(cambios),
                D.excel_tablas([("Cambios propuestos", cambios)]),
                "Gemweb_cambios_propuestos_%s.xlsx" % fecha.strftime("%Y%m%d_%H%M"), XLSX,
                key="sgc_gw_excel_cambios_%s" % p["fecha"])
            editor = cambios.copy()
            editor.insert(0, "Aplicar", False)
            elegidos = st.data_editor(
                editor, hide_index=True, use_container_width=True, height=320,
                disabled=list(cambios.columns), key="sgc_gw_sel_%s" % p["fecha"],
                column_config={"Aplicar": st.column_config.CheckboxColumn("Aplicar")})
            marcados = [c for c, sel in zip(p["cambios"], elegidos["Aplicar"]) if sel]
            b1, b2 = st.columns(2)
            todos = b1.button("Aplicar todos (%d)" % len(p["cambios"]), type="primary",
                              icon=":material/done_all:", key="sgc_gw_todos",
                              use_container_width=True)
            seleccion = b2.button("Aplicar los seleccionados (%d)" % len(marcados),
                                  icon=":material/done:", key="sgc_gw_sel",
                                  disabled=not marcados, use_container_width=True)
            if todos or seleccion:
                hechos, saltados = G.aplicar(p["cambios"] if todos else marcados, _usuario())
                ss.sgc_gw_aviso = ("Aplicados %d cambios al cuadro." % len(hechos)
                                   + (" %d no, porque alguien había cambiado esa celda después "
                                      "de la búsqueda." % len(saltados) if saltados else ""))
                ss.sgc_gw_abierto = True
                st.rerun()
        if ss.get("sgc_gw_aviso"):
            st.success(ss.pop("sgc_gw_aviso"))
    with t2:
        ver = st.multiselect("Resultado", sorted(listado["Resultado"].dropna().unique()),
                             key="sgc_gw_ver")
        st.dataframe(listado[listado["Resultado"].isin(ver)] if ver else listado,
                     hide_index=True, use_container_width=True, height=350,
                     column_config=FECHAS_LISTADO)


# ------------------------------------------------------------------- retraso
def _grafico_barras(tabla, por, color):
    return alt.Chart(tabla).mark_bar(color=color).encode(
        x=alt.X("% Retraso:Q", axis=alt.Axis(format="%", title="% de retraso"),
                scale=alt.Scale(domain=[0, 1])),
        y=alt.Y("%s:N" % por, sort="-x", title=None),
        tooltip=[por, "Retrasos", "Enviados",
                 alt.Tooltip("% Retraso:Q", format=".1%")]).properties(
        height=max(120, 32 * len(tabla)))


def excel_retraso(df, tec, aux, indicadores, hoy):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    r, e, p = indicadores
    ws.append(["TOTAL RETRASOS", r, "TOTAL ENVIADOS", e, "% CUMPLIMIENTO", p, None,
               "Fecha de control:", hoy])
    ws.append([])
    ws.append(["Técnico", "Retrasos", "Enviados", "% Retraso", None,
               "Auxiliar", "Retrasos", "Enviados", "% Retraso"])
    for i in range(max(len(tec), len(aux))):
        a = list(tec.iloc[i]) if i < len(tec) else [None] * 4
        b = list(aux.iloc[i]) if i < len(aux) else [None] * 4
        ws.append(a + [None] + b)
    for fila in ws.iter_rows(min_row=4):
        for x in (fila[3], fila[8]):
            x.number_format = "0.00%"
    ws["F1"].number_format = "0.00%"
    ws["I1"].number_format = "dd/mm/yyyy"
    salida = io.BytesIO()
    wb.save(salida)
    auditoria = df.drop(columns=["_id"])
    wb2 = openpyxl.load_workbook(io.BytesIO(D.excel_tablas([("Auditoría", auditoria)])))
    destino = wb.create_sheet("Auditoría")
    for fila in wb2["Auditoría"].iter_rows():
        destino.append([x.value for x in fila])
    for col in ("H", "I"):
        for x in destino[col][1:]:
            x.number_format = "mmm-yy"
    for x in destino["J"][1:]:
        x.number_format = "dd/mm/yyyy"
    destino.freeze_panes = "A2"
    destino.auto_filter.ref = destino.dimensions
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def pagina_retraso():
    if not _hay_datos():
        return
    c1, c2 = st.columns([1, 1])
    with c1:
        anio = _anio()
    hoy = c2.date_input("Fecha de control", dt.date.today(), format="DD/MM/YYYY",
                        key="sgc_ret_hoy",
                        help="Se evalúan los bloques cuya fecha de control (día 1 del mes M+2) "
                             "ya ha llegado en esta fecha.")
    suministros = almacen.listar_sgc_suministros(bajas=False)
    df = R.auditoria_df(R.auditoria(suministros, almacen.listar_sgc_estados(anio), anio, hoy))
    st.caption("Solo los suministros «Sgto. y Control». Cada bloque (1, 3, 6 o 12 meses según la "
               "cadencia) es ENVIADO si todos sus meses con factura tienen «Inf enviado» y RETRASO "
               "si no. «Sin fra», «Cadencia», «Temporal» y los vacíos no esperan informe.")
    if df.empty:
        st.info("No hay bloques que evaluar con esa fecha de control.")
        return

    f = st.columns(4)
    tec = f[0].multiselect("Técnico", sorted(df["Técnico"].replace("", pd.NA).dropna().unique()),
                           key="sgc_ret_tec")
    aux = f[1].multiselect("Auxiliar", sorted(df["Auxiliar"].replace("", pd.NA).dropna().unique()),
                           key="sgc_ret_aux")
    cli = f[2].multiselect("Cliente/Grupo", sorted(df["Cliente/Grupo"].dropna().unique()),
                           key="sgc_ret_cli")
    texto = f[3].text_input("Buscar (suministro, CUPS, observaciones)",
                            key="sgc_ret_txt").strip().lower()
    g = st.columns(4)
    cad = g[0].multiselect("Cadencia", sorted(df["Cadencia"].dropna().unique()), key="sgc_ret_cad")
    res = g[1].multiselect("Resultado", [R.RETRASO, R.ENVIADO], key="sgc_ret_res")
    meses = sorted(df["Mes fin"].unique())
    etiqueta = lambda d: "%s-%02d" % (C.MESES_CORTOS[d.month - 1], d.year % 100)
    desde = g[2].selectbox("Bloques que terminan desde", meses, format_func=etiqueta,
                           key="sgc_ret_desde")
    hasta = g[3].selectbox("hasta", meses, index=len(meses) - 1, format_func=etiqueta,
                           key="sgc_ret_hasta")
    if tec:
        df = df[df["Técnico"].isin(tec)]
    if aux:
        df = df[df["Auxiliar"].isin(aux)]
    if cli:
        df = df[df["Cliente/Grupo"].isin(cli)]
    if cad:
        df = df[df["Cadencia"].isin(cad)]
    if res:
        df = df[df["Resultado"].isin(res)]
    df = df[(df["Mes fin"] >= desde) & (df["Mes fin"] <= hasta)]
    if texto:
        df = df[df[["Denominación suministro", "CUPS", "Observaciones"]].fillna("").apply(
            lambda c: c.str.lower().str.contains(texto, regex=False)).any(axis=1)]

    indic = R.indicadores(df)
    m = st.columns(4)
    m[0].metric("Retrasos", indic[0])
    m[1].metric("Enviados", indic[1])
    m[2].metric("Cumplimiento", "—" if indic[2] is None else
                ("%.2f %%" % (indic[2] * 100)).replace(".", ","))
    m[3].metric("Bloques evaluados", len(df))
    sin_tecnico = int((df["Técnico"] == "").sum())
    if sin_tecnico:
        st.warning("%d bloques son de suministros sin técnico: cuentan en el total pero no en la "
                   "tabla por técnico." % sin_tecnico)

    tabla_tec, tabla_aux = R.resumen(df, "Técnico"), R.resumen(df, "Auxiliar")
    c1, c2 = st.columns(2)
    for col, tabla, por, color in ((c1, tabla_tec, "Técnico", "#0078D7"),
                                   (c2, tabla_aux, "Auxiliar", "#FF5A5A")):
        with col:
            st.markdown("**%% de retraso por %s**" % por.lower())
            if tabla.empty:
                st.caption("Sin datos.")
                continue
            st.altair_chart(_grafico_barras(tabla, por, color), use_container_width=True)
            st.dataframe(tabla, hide_index=True, use_container_width=True, column_config={
                "% Retraso": st.column_config.NumberColumn(format="percent")})

    st.markdown("**Auditoría** (un bloque por fila; pulsa una fila para ver el suministro)")
    estilo_res = lambda v: ("background-color:#FFC7CE;color:#9C0006;font-weight:bold"
                            if v == R.RETRASO else
                            "background-color:#C6EFCE;color:#006100;font-weight:bold"
                            if v == R.ENVIADO else "")
    sel = st.dataframe(
        df.drop(columns=["_id"]).style.map(estilo_res, subset=["Resultado"]), hide_index=True,
        use_container_width=True, height=450, on_select="rerun", selection_mode="single-row",
        key="sgc_ret_tabla", column_config={
            "Mes inicio": st.column_config.DateColumn(format="MMM YYYY"),
            "Mes fin": st.column_config.DateColumn(format="MMM YYYY"),
            "Fecha control (M+2)": st.column_config.DateColumn(format="DD/MM/YYYY")})
    filas = sel.selection.rows if sel else []
    if filas:
        id_ = int(df.iloc[filas[0]]["_id"])
        cuadro = D.cuadro_df(anio)
        fila = cuadro[cuadro[D.ID] == id_]
        if len(fila):
            st.markdown("**Suministro en el cuadro (%d)**" % anio)
            mostrar = ["Técnico", "Auxiliar", "Cliente/Grupo", "Denominación suministro", "CUPS",
                       "Cadencia informe"] + [c for c in D.columnas_cuadro(anio)
                                              if c[:3] in C.MESES_CORTOS or c.startswith("Obs ")]
            st.dataframe(fila[mostrar], hide_index=True, use_container_width=True)
    _boton_excel("Retraso_informes_%d.xlsx" % anio,
                 lambda: excel_retraso(df, tabla_tec, tabla_aux, indic, hoy), "sgc_ret_excel")


# ------------------------------------------------------------------- tablas auxiliares
def _pagina_tabla(tabla, ayuda=""):
    if not _hay_datos():
        return
    t = C.TABLAS[tabla]
    listas, equipo = almacen.listar_sgc_listas(), almacen.listar_sgc_equipo()
    df = D.tabla_df(tabla)
    if ayuda:
        st.caption(ayuda)
    df, firma = _filtrar(df, "sgc_" + tabla)
    st.metric("Filas", len(df))
    editado = st.data_editor(
        df, hide_index=True, use_container_width=True, num_rows="dynamic", height=600,
        column_config=_config_columnas(t["columnas"], df, listas, equipo),
        key="sgc_ed_%s_%s" % (tabla, _firma(firma)))
    if st.button("Guardar cambios", type="primary", icon=":material/save:",
                 key="sgc_guardar_" + tabla):
        n, b = D.guardar_tabla(tabla, df, editado, _usuario())
        almacen.registrar_actividad(_usuario(), MODULO, "Guardar %s" % t["titulo"],
                                    "%d filas, %d borradas" % (n, b))
        st.success("Guardado: %d filas cambiadas o nuevas y %d borradas." % (n, b))
        st.rerun()
    _boton_excel("%s.xlsx" % t["titulo"].replace(" ", "_"),
                 lambda: D.excel_tablas([(t["titulo"], D.tabla_df(tabla))]), "sgc_x_" + tabla)
    return editado


def emails_invalidos(df):
    """Filas con algun email mal escrito en Para o CC."""
    malas = []
    for _, r in df.iterrows():
        for col in ("Para", "CC"):
            for e in re.split(r"[;,]", str(r.get(col) or "")):
                e = e.strip()
                if e and not EMAIL.fullmatch(e):
                    malas.append({"Cliente/Grupo": r.get("Cliente/Grupo"),
                                  "Denominación suministro": r.get("Denominación suministro"),
                                  "Columna": col, "Email": e})
    return pd.DataFrame(malas)


def sin_contacto_envio():
    """Suministros activos «Sgto. y Control» del cuadro sin contacto de envio (por CUPS)."""
    cups = {(f["datos"].get("CUPS") or "").strip().upper()[:20]
            for f in almacen.listar_sgc_filas("contactos_envio")}
    filas = []
    for s in almacen.listar_sgc_suministros(bajas=False):
        d = s["datos"]
        if (d.get("Tipo contratación") or "").strip() != C.SGTO_Y_CONTROL:
            continue
        if (d.get("CUPS") or "").strip().upper()[:20] not in cups:
            filas.append({k: d.get(k) for k in ("Técnico", "Auxiliar", "Cliente/Grupo",
                                                "Denominación suministro", "CUPS")})
    return pd.DataFrame(filas)


def pagina_contactos_envio():
    editado = _pagina_tabla(
        "contactos_envio", "A quién se envían los informes de seguimiento y los avisos de "
        "incidencias: una fila por suministro. «Nombre cliente» es el del saludo del correo; "
        "en «Para» y «CC» se separan varios emails con «;».")
    if editado is None:
        return
    malas = emails_invalidos(editado)
    if len(malas):
        with st.expander(":material/warning: %d emails mal escritos" % len(malas)):
            st.dataframe(malas, hide_index=True, use_container_width=True)
    sin = sin_contacto_envio()
    if len(sin):
        with st.expander(":material/warning: %d suministros «Sgto. y Control» sin contacto de "
                         "envío" % len(sin)):
            st.caption("Se buscan por CUPS. No recibirán informes ni avisos hasta que tengan su "
                       "fila aquí.")
            st.dataframe(sin, hide_index=True, use_container_width=True)


def pagina_contactos():
    _pagina_tabla("contactos", "Personas de contacto de cada cliente.")


def pagina_claves():
    _pagina_tabla("claves", "Accesos de cada cliente a los portales de su comercializadora.")


def pagina_smac():
    _pagina_tabla("smac")


def pagina_telemedida():
    _pagina_tabla("telemedida")


# ------------------------------------------------------------------- bajas
def pagina_bajas():
    if not _hay_datos():
        return
    anio = dt.date.today().year
    df = D.cuadro_df(anio, bajas=True)
    if df.empty:
        st.info("No hay suministros dados de baja.")
        return
    df, _ = _filtrar(df, "sgc_bajas")
    columnas = ["Fecha de baja", "Técnico", "Auxiliar", "Cliente/Grupo", "Titular de contrato",
                "CUPS", "Denominación suministro", "Tipo contratación", "Cadencia informe",
                "Comercializadora", "Fecha fin contrato", C.TOTAL_CUOTA, "Observaciones", C.OTROS]
    st.metric("Suministros de baja", len(df))
    sel = st.dataframe(df[[D.ID] + columnas], hide_index=True, use_container_width=True,
                       height=500, on_select="rerun", selection_mode="multi-row",
                       key="sgc_bajas_tabla", column_config={
                           D.ID: None,
                           "Fecha de baja": st.column_config.DateColumn(format="DD/MM/YYYY"),
                           "Fecha fin contrato": st.column_config.DateColumn(format="DD/MM/YYYY"),
                           C.TOTAL_CUOTA: st.column_config.NumberColumn(format="localized")})
    filas = sel.selection.rows if sel else []
    if filas and st.button("Reactivar %d suministro(s)" % len(filas), icon=":material/undo:"):
        D.dar_de_baja([int(df.iloc[i][D.ID]) for i in filas], _usuario(), baja=False)
        st.success("Reactivados: vuelven al cuadro de seguimiento.")
        st.rerun()


# ------------------------------------------------------------------- administracion
def pagina_admin():
    if not ss.usuario["admin"]:
        st.warning("Solo los administradores pueden importar y configurar Seguimiento y Control.")
        return
    pestanas = st.tabs(["Importar Excel", "Equipo", "Listas de valores", "Outlook"])
    with pestanas[0]:
        _admin_importar()
    with pestanas[1]:
        _admin_equipo()
    with pestanas[2]:
        _admin_listas()
    with pestanas[3]:
        from seguimiento import envios
        envios.admin_outlook()


def _admin_importar():
    st.caption("Pasa a la plataforma el Excel combinado de cuadros de seguimiento (pestañas "
               "CLIENTES, BAJAS, CLAVES, CONTACTOS, SMAC, TELEMEDIDA y Lista) y el Excel de "
               "contactos de envío. **Sustituye todo** lo que haya en Seguimiento y Control; "
               "antes se hace una copia de la base de datos. Los ficheros no se guardan: solo "
               "su contenido, en la base de datos.")
    c1, c2 = st.columns(2)
    cuadro = c1.file_uploader("Excel del cuadro de seguimiento", type=["xlsx", "xlsm"],
                              key="sgc_imp_cuadro")
    contactos = c2.file_uploader("Excel de contactos de envío (opcional: si no se sube, se "
                                 "conservan los que ya hay)", type=["xlsx", "xlsm"],
                                 key="sgc_imp_contactos")
    if not cuadro:
        return
    try:
        leido = D.leer_cuadro(cuadro.getvalue())
        lista_contactos = D.leer_contactos_envio(contactos.getvalue()) if contactos else None
    except Exception as e:                                          # noqa: BLE001
        st.error("No se ha podido leer: %s" % e)
        return
    resumen = [{"Pestaña": k, "Filas": v} for k, v in leido["recuento"].items()]
    resumen.append({"Pestaña": "Estados de meses", "Filas": len(leido["estados"])})
    if lista_contactos is not None:
        resumen.append({"Pestaña": "Contactos de envío", "Filas": len(lista_contactos)})
    st.dataframe(pd.DataFrame(resumen), hide_index=True)
    con_otros = sum(1 for s in leido["suministros"] if s["datos"].get(C.OTROS))
    if con_otros:
        st.info("%d suministros tienen valores que no encajan en su columna (fechas o números "
                "escritos como texto, columnas sin título): se guardan en la columna «Otros» "
                "con su título para que no se pierdan." % con_otros)
    if st.button("Importar y sustituir", type="primary", icon=":material/upload:"):
        copia = almacen.copia_seguridad(config.CARPETA_DATOS / "copias_seguridad")
        D.importar(leido, lista_contactos, _usuario())
        almacen.registrar_actividad(_usuario(), MODULO, "Importar Excel", cuadro.name)
        st.success("Importado. Copia previa de la base de datos: %s" % copia.name)


def _admin_equipo():
    st.caption("Técnicos y auxiliares de los desplegables del cuadro. El email del técnico va en "
               "copia en los correos; el usuario de la plataforma sirve para preseleccionar el "
               "auxiliar de quien envía.")
    usuarios = [""] + [u["usuario"] for u in almacen.listar_usuarios()]
    df = pd.DataFrame(almacen.listar_sgc_equipo(),
                      columns=["nombre", "tecnico", "auxiliar", "email", "usuario"])
    df["tecnico"] = df["tecnico"].astype(bool)
    df["auxiliar"] = df["auxiliar"].astype(bool)
    editado = st.data_editor(df, hide_index=True, num_rows="dynamic", use_container_width=True,
                             key="sgc_equipo", column_config={
                                 "nombre": st.column_config.TextColumn("Nombre", required=True),
                                 "tecnico": st.column_config.CheckboxColumn("Técnico"),
                                 "auxiliar": st.column_config.CheckboxColumn("Auxiliar"),
                                 "email": st.column_config.TextColumn("Email"),
                                 "usuario": st.column_config.SelectboxColumn(
                                     "Usuario de la plataforma", options=usuarios)})
    if st.button("Guardar equipo", type="primary", key="sgc_guardar_equipo"):
        try:
            almacen.guardar_sgc_equipo(editado.fillna("").to_dict("records"))
        except ValueError as e:
            st.error(str(e))
        else:
            st.success("Equipo guardado.")
            st.rerun()


def _admin_listas():
    listas = almacen.listar_sgc_listas()
    nombres = [n for n in C.TITULOS_LISTAS if n not in ("tecnicos", "auxiliares")]
    lista = st.selectbox("Lista", nombres, format_func=C.TITULOS_LISTAS.get, key="sgc_lista")
    df = pd.DataFrame(listas.get(lista, []), columns=["valor", "descripcion", "color"])
    cfg = {"valor": st.column_config.TextColumn("Valor", required=True),
           "descripcion": st.column_config.TextColumn("Descripción"),
           "color": st.column_config.TextColumn("Color (#RRGGBB)")}
    if lista != "estados":
        df = df.drop(columns=["color"])
        cfg.pop("color")
    else:
        st.caption("Los estados «Sin fra», «Cadencia» y «Temporal» no esperan informe en el "
                   "cálculo del retraso; «Inf enviado» es el único que cuenta como enviado.")
    editado = st.data_editor(df, hide_index=True, num_rows="dynamic", use_container_width=True,
                             column_config=cfg, key="sgc_lista_ed_" + lista)
    if st.button("Guardar lista", type="primary", key="sgc_guardar_lista"):
        valores = [{k: (None if pd.isna(v) or v == "" else str(v).strip()) for k, v in r.items()}
                   for r in editado.to_dict("records")]
        almacen.guardar_sgc_lista(lista, [v for v in valores if v.get("valor")])
        st.success("Lista guardada.")
        st.rerun()
