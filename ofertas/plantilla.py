# -*- coding: utf-8 -*-
"""
PLANTILLA DE SUMINISTROS Y CONSUMOS PARA LA SOLICITUD DE OFERTAS
================================================================
Para pedir ofertas con otros consumos que los de las facturas de Gemweb, o de suministros que
no estan en Gemweb. Se descarga un Excel con tres pestañas (Instrucciones, Suministros y
Consumos), prellenado con lo que ha encontrado Gemweb; se cambia lo necesario o se añaden
suministros y se vuelve a subir:

- Suministros: una fila por CUPS (datos del cliente, tarifa y potencias P1-P6; en gas, Qd).
- Consumos: 12 filas por CUPS (enero a diciembre) con el consumo de P1 a P6 (en gas, el kWh
  del mes).

aplicar_plantilla() devuelve las fichas de la solicitud con lo de la plantilla: en los CUPS de
Gemweb manda lo que este relleno en la plantilla; los CUPS nuevos se crean como suministros
manuales. Las fichas tienen las mismas claves que gemweb_extractor.paginas.ficha_solicitud,
asi que el Excel de la solicitud, el correo y el guardado no cambian. Sin Streamlit.
"""

import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ofertas import solicitud as so

GEMWEB, MANUAL, PLANTILLA = "Gemweb", "Manual", so.CONSUMO_PLANTILLA
HOJA_SUMINISTROS, HOJA_CONSUMOS, HOJA_INSTRUCCIONES = "Suministros", "Consumos", "Instrucciones"
FILAS_EN_BLANCO = 10                     # filas vacias en Suministros para añadir suministros

# columnas de la hoja Suministros: (clave de la ficha, titulo)
DATOS = [("cups", "CUPS"), ("grupo", "Cliente/Grupo"), ("razon", "Razón social"),
         ("cif", "CIF"), ("alias", "Alias"), ("direccion", "Dirección de suministro"),
         ("tarifa", "Tarifa")]
POTENCIAS = ["Potencia %s (kW)" % p for p in so.PERIODOS]
QD = "Qd (kWh/día)"
ORIGEN = "Origen"
CONSUMO_P = ["%s kWh" % p for p in so.PERIODOS]
CONSUMO_GAS = "Consumo (kWh)"
MES, NOMBRE_MES = "Mes", "Nombre del mes"
NOMBRES_MES = [m.capitalize() for m in so.MESES_LARGOS]
TARIFA_ELECTRICA = re.compile(r"^(2\.0|3\.0|6\.[1-4])\s*TD", re.I)

GRANATE, ROSA, GRIS = "970000", "F3D9D9", "F2F2F2"


def _columnas_suministros(gas):
    return [t for _, t in DATOS] + ([QD] if gas else POTENCIAS) + [ORIGEN]


def _columnas_consumos(gas):
    return ["CUPS", MES, NOMBRE_MES] + ([CONSUMO_GAS] if gas else CONSUMO_P + ["Total kWh"])


def _cabecera(ws, titulos, anchos):
    ws.append(titulos)
    for n, celda in enumerate(ws[1], 1):
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=GRANATE)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(n)].width = anchos.get(celda.value, 14)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"


def plantilla_excel(fichas, cups_sin_ficha=(), producto=so.ELECTRICIDAD) -> bytes:
    """Excel para rellenar: con los suministros de Gemweb (fichas) y sus consumos, los CUPS que
    Gemweb no ha encontrado (solo el CUPS) y filas vacias para añadir otros."""
    gas = producto == so.GAS
    wb = Workbook()
    ins = wb.active
    ins.title = HOJA_INSTRUCCIONES
    ins.column_dimensions["A"].width = 110
    lineas = [
        "Plantilla para la solicitud de ofertas (%s)" % ("gas natural" if gas else
                                                          "electricidad"),
        "",
        "1. Pestaña «Suministros»: una fila por CUPS. Los de Gemweb ya vienen rellenos (Origen "
        "= Gemweb): cambia solo lo que quieras; lo que dejes en blanco se queda como en Gemweb.",
        "   Para un suministro que no está en Gemweb, rellena una fila nueva: CUPS, razón "
        "social, tarifa y %s son obligatorios." % ("Qd" if gas else "potencias P1-P6"),
        "2. Pestaña «Consumos»: 12 filas por CUPS (meses 1 a 12, de enero a diciembre) con el "
        + ("consumo del mes en kWh." if gas else "consumo de cada periodo en kWh (P1 a P6)."),
        "   Los de Gemweb vienen con el perfil de sus facturas: cambia los meses que quieras. "
        "Si un CUPS de Gemweb no tiene filas aquí, se queda con el consumo de Gemweb.",
        "   Para un suministro nuevo, añade sus 12 filas (copia el bloque de otro y cambia el "
        "CUPS y los consumos).",
        "3. Guarda el Excel y súbelo en la página de la solicitud. Antes de aplicarlo verás los "
        "cambios y, si falta algo, en qué pestaña y fila.",
        "",
        "Números con coma o punto decimal. La columna «Total kWh» y «Nombre del mes» son solo "
        "informativas.",
    ]
    for texto in lineas:
        ins.append([texto])
    ins["A1"].font = Font(bold=True, size=13, color=GRANATE)

    ws = wb.create_sheet(HOJA_SUMINISTROS)
    anchos = {"CUPS": 25, "Cliente/Grupo": 24, "Razón social": 30, "Dirección de suministro": 40,
              "Alias": 18, "CIF": 12, "Tarifa": 10, ORIGEN: 10}
    _cabecera(ws, _columnas_suministros(gas), anchos)
    for f in fichas:
        fila = [f.get(k) or "" for k, _ in DATOS]
        if gas:
            fila.append((f.get("tecnicos") or {}).get("qd"))
        else:
            fila += list(f.get("potencias") or [None] * 6)
        ws.append(fila + [f.get("origen") or GEMWEB])
    for cups in cups_sin_ficha:
        ws.append([cups] + [None] * (len(_columnas_suministros(gas)) - 2) + [MANUAL])
    for _ in range(FILAS_EN_BLANCO):
        ws.append([None] * (len(_columnas_suministros(gas)) - 1) + [MANUAL])
    for fila in ws.iter_rows(min_row=2):
        fila[-1].fill = PatternFill("solid", fgColor=GRIS)

    wc = wb.create_sheet(HOJA_CONSUMOS)
    _cabecera(wc, _columnas_consumos(gas), {"CUPS": 25, NOMBRE_MES: 14})
    for f in fichas:
        for m in range(1, 13):
            if gas:
                p = (f.get("perfil_gas") or {}).get(m) or {}
                wc.append([f["cups"], m, NOMBRES_MES[m - 1], round(p.get("kwh") or 0)])
            else:
                valores = [round(v) for v in ((f.get("perfil") or {}).get(m) or [0] * 6)]
                r = wc.max_row + 1
                wc.append([f["cups"], m, NOMBRES_MES[m - 1]] + valores
                          + ["=SUM(D%d:I%d)" % (r, r)])
    for cups in cups_sin_ficha:
        for m in range(1, 13):
            r = wc.max_row + 1
            wc.append([cups, m, NOMBRES_MES[m - 1]] + ([None] if gas else [None] * 6
                                                        + ["=SUM(D%d:I%d)" % (r, r)]))
    for fila in wc.iter_rows(min_row=2):
        for c in fila[3:]:
            c.number_format = "#,##0"
    wb.active = 1
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ------------------------------------------------------------------- lectura
def _cups(v) -> str:
    return re.sub(r"\s", "", str(v or "")).upper()


def _numero(v):
    """Numero de una celda (acepta coma decimal y separador de miles con punto). None si esta
    vacia; ValueError si no es un numero."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "")
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    return float(t)


def _filas(ws):
    """[(numero de fila, {titulo: valor})] de una hoja con la cabecera en la fila 1."""
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    titulos = [str(t).strip() if t is not None else "" for t in filas[0]]
    return [(n, dict(zip(titulos, f))) for n, f in enumerate(filas[1:], 2)]


def leer_plantilla(contenido: bytes, producto=so.ELECTRICIDAD):
    """-> (suministros {cups: datos}, consumos {cups: {mes: valores}}, errores [texto]).
    datos: claves de la ficha que vienen rellenas (y potencias o qd); valores: [P1..P6] en
    electricidad o kWh en gas."""
    gas = producto == so.GAS
    errores = []
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True)
    except Exception:                                     # noqa: BLE001
        return {}, {}, ["No se puede abrir el fichero: tiene que ser la plantilla en Excel "
                        "(.xlsx)."]
    for hoja in (HOJA_SUMINISTROS, HOJA_CONSUMOS):
        if hoja not in wb.sheetnames:
            errores.append("Falta la pestaña «%s»." % hoja)
    if errores:
        return {}, {}, errores

    suministros = {}
    for n, f in _filas(wb[HOJA_SUMINISTROS]):
        cups = _cups(f.get("CUPS"))
        rellenas = [k for k, v in f.items() if k not in (ORIGEN, "CUPS") and v not in (None, "")]
        if not cups:
            if rellenas:
                errores.append("Suministros, fila %d: falta el CUPS." % n)
            continue
        if not rellenas and str(f.get(ORIGEN) or MANUAL).strip() == MANUAL:
            continue                       # CUPS no encontrado que se ha dejado sin rellenar
        if not (cups.startswith("ES") and 20 <= len(cups) <= 22):
            errores.append("Suministros, fila %d: el CUPS «%s» no es válido." % (n, cups))
            continue
        if cups in suministros:
            errores.append("Suministros, fila %d: el CUPS %s está repetido." % (n, cups))
            continue
        datos = {"origen": str(f.get(ORIGEN) or MANUAL).strip()}
        for clave, titulo in DATOS[1:]:
            valor = f.get(titulo)
            if valor not in (None, ""):
                datos[clave] = so.cif_limpio(valor) if clave == "cif" else str(valor).strip()
        try:
            if gas:
                if _numero(f.get(QD)) is not None:
                    datos["qd"] = _numero(f.get(QD))
            else:
                potencias = [_numero(f.get(t)) for t in POTENCIAS]
                if any(p is not None for p in potencias):
                    datos["potencias"] = potencias
        except ValueError:
            errores.append("Suministros, fila %d (%s): %s no es un número." % (
                n, cups, "el Qd" if gas else "alguna potencia"))
            continue
        if any(v is not None and v < 0 for v in datos.get("potencias", []) + [datos.get("qd")]):
            errores.append("Suministros, fila %d (%s): hay valores negativos." % (n, cups))
        suministros[cups] = datos

    consumos, rellenos = {}, set()
    columnas = [CONSUMO_GAS] if gas else CONSUMO_P
    for n, f in _filas(wb[HOJA_CONSUMOS]):
        cups = _cups(f.get("CUPS"))
        if not cups:
            continue
        try:
            mes = int(_numero(f.get(MES)) or 0)
            valores = [_numero(f.get(c)) for c in columnas]
        except ValueError:
            errores.append("Consumos, fila %d (%s): el mes o algún consumo no es un número."
                           % (n, cups))
            continue
        if not 1 <= mes <= 12:
            errores.append("Consumos, fila %d (%s): el mes tiene que ser de 1 a 12." % (n, cups))
            continue
        if any(v is not None and v < 0 for v in valores):
            errores.append("Consumos, fila %d (%s): hay consumos negativos." % (n, cups))
            continue
        meses = consumos.setdefault(cups, {})
        if mes in meses:
            errores.append("Consumos, fila %d (%s): el mes %d está repetido." % (n, cups, mes))
            continue
        if any(v is not None for v in valores):
            rellenos.add(cups)
        valores = [v or 0.0 for v in valores]          # un mes en blanco cuenta como 0
        meses[mes] = valores[0] if gas else valores
    # los CUPS con todas sus filas en blanco (no encontrados que no se rellenan) no cuentan
    consumos = {c: m for c, m in consumos.items() if c in rellenos}
    for cups, meses in consumos.items():
        faltan = [m for m in range(1, 13) if m not in meses]
        if faltan:
            errores.append("Consumos, %s: faltan los meses %s." % (
                cups, ", ".join(map(str, faltan))))
        if cups not in suministros:
            errores.append("Consumos, %s: no está en la pestaña «Suministros»." % cups)
    return suministros, consumos, errores


# ------------------------------------------------------------------- aplicar
def _consumo_vacio(meses, gas) -> bool:
    return not any((v if gas else sum(v)) for v in meses.values())


def _ficha_manual(cups, datos, gas):
    ficha = {"cups": cups, "alias": datos.get("alias") or "", "tipo": so.GAS if gas
             else so.ELECTRICIDAD, "grupo": datos.get("grupo") or "",
             "razon": datos.get("razon") or "", "cif": datos.get("cif") or "",
             "direccion": datos.get("direccion") or "", "tarifa": datos.get("tarifa") or "",
             "fecha_inicio": "", "potencias": datos.get("potencias") or [None] * 6,
             "perfil": {m: [0.0] * 6 for m in range(1, 13)}, "aviso": "", "aviso_grave": False,
             "facturas": 0, "ultima_factura": None, "origen": MANUAL}
    if gas:
        ficha["perfil_gas"] = {m: {"anio": None, "kwh": 0.0, "dias": 0, "qd": None}
                               for m in range(1, 13)}
    return ficha


def errores_de_manuales(suministros, consumos, fichas, gas) -> list:
    """Lo obligatorio de los suministros que no vienen de Gemweb."""
    de_gemweb = {f["cups"] for f in fichas if f.get("origen") != MANUAL}
    errores = []
    for cups, d in suministros.items():
        if cups in de_gemweb:
            continue
        faltan = []
        if not d.get("razon"):
            faltan.append("razón social")
        tarifa = d.get("tarifa") or ""
        if not tarifa:
            faltan.append("tarifa")
        elif gas and not so.grupo_rl(tarifa):
            errores.append("Suministros, %s: la tarifa «%s» no es un peaje de gas (RL.1 a "
                           "RL.11)." % (cups, tarifa))
        elif not gas and not TARIFA_ELECTRICA.match(tarifa):
            errores.append("Suministros, %s: la tarifa «%s» no es una tarifa eléctrica (2.0TD, "
                           "3.0TD, 6.1TD…)." % (cups, tarifa))
        if gas and d.get("qd") is None:
            faltan.append("Qd")
        if not gas and not any(p is not None for p in d.get("potencias") or []):
            faltan.append("potencias")
        if cups not in consumos or _consumo_vacio(consumos[cups], gas):
            faltan.append("consumos (pestaña «Consumos»)")
        if faltan:
            errores.append("Suministros, %s (no está en Gemweb): falta %s." % (
                cups, ", ".join(faltan)))
    return errores


def aplicar_plantilla(fichas, leido, producto=so.ELECTRICIDAD):
    """(fichas nuevas, cambios). leido: lo que devuelve leer_plantilla (sin errores).
    cambios: una fila por suministro con su consumo anual antes y despues."""
    gas = producto == so.GAS
    suministros, consumos, _ = leido
    nuevas, cambios = [], []
    por_cups = {f["cups"]: f for f in fichas}
    orden = list(por_cups) + [c for c in suministros if c not in por_cups]
    for cups in orden:
        anterior = por_cups.get(cups)
        datos = suministros.get(cups, {})
        f = dict(anterior) if anterior else _ficha_manual(cups, datos, gas)
        antes = so.consumo_anual(anterior) if anterior else None
        for clave in ("grupo", "razon", "cif", "alias", "direccion", "tarifa"):
            if datos.get(clave):
                f[clave] = datos[clave]
        if datos.get("potencias") and not gas:
            f["potencias"] = [p if p is not None else viejo for p, viejo in zip(
                datos["potencias"], f.get("potencias") or [None] * 6)]
        meses = consumos.get(cups)
        if meses and not (anterior and _consumo_vacio(meses, gas)):
            if gas:
                previo = f.get("perfil_gas") or {}
                cambia = any(round(meses[m]) != round((previo.get(m) or {}).get("kwh") or 0)
                             for m in range(1, 13))
            else:
                previo = f.get("perfil") or {}
                cambia = any(round(v) != round(w) for m in range(1, 13)
                             for v, w in zip(meses[m], previo.get(m) or [0] * 6))
            # la plantilla sale con el consumo de Gemweb redondeado: si no se ha tocado, se
            # queda el de Gemweb
            if cambia or not anterior:
                if gas:
                    perfil = {m: dict(previo.get(m) or {"anio": None, "dias": 0, "qd": None})
                              for m in range(1, 13)}
                    for m, kwh in meses.items():
                        perfil[m]["kwh"] = kwh
                    f["perfil_gas"] = perfil
                else:
                    f["perfil"] = {m: list(meses[m]) for m in range(1, 13)}
                f["origen_consumo"] = PLANTILLA
                f["aviso"], f["aviso_grave"] = "", False
        if gas:
            # lo que depende del consumo y del Qd se recalcula; lo demas se conserva
            qd = datos.get("qd", (f.get("tecnicos") or {}).get("qd"))
            previos = {k: v for k, v in (f.get("tecnicos") or {}).items()
                       if v is not None and k not in ("qd", "telemedida", "termino_fijo")}
            f["tecnicos"] = dict(so.tecnicos_por_defecto(so.consumo_anual(f), qd), **previos)
        despues = so.consumo_anual(f)
        cambios.append({"CUPS": cups, "Alias": f.get("alias") or "",
                        "Origen": f.get("origen") or GEMWEB,
                        "Consumo Gemweb (kWh)": round(antes) if antes is not None else None,
                        "Consumo nuevo (kWh)": round(despues),
                        "Diferencia (%)": round((despues - antes) / antes * 100, 1)
                        if antes else None})
        nuevas.append(f)
    return nuevas, cambios


def origen_consumo(f) -> str:
    return f.get("origen_consumo") or (MANUAL if f.get("origen") == MANUAL else GEMWEB)
