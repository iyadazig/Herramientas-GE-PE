# -*- coding: utf-8 -*-
"""
SOLICITUD DE OFERTAS (Excel con el formato de GE&PE)
====================================================
Hoja «Resumen» con un suministro por fila y una hoja por suministro con la cabecera de la
solicitud, los datos del suministro, las potencias de la ultima factura y el perfil de
consumo de los ultimos 12 meses por periodos (cada factura mensual en el mes con mas dias
suyos; las bimestrales o mas largas, repartidas por dias: reparto_mensual), con su grafico
de columnas apiladas.

En gas natural (cabecera producto_tipo = GAS), en lugar de potencias y periodos: los aspectos
tecnicos del suministro (peaje RL, Qd, telemedida, termino fijo...) y el perfil mensual con
el Qd medio y el Qd aplicado, como en las peticiones de gas de GE&PE. Las reglas de la
normativa de gas (Circular CNMC 6/2020, Orden IET/2446/2013, Ley 38/1992) estan aqui.

Los datos salen de Gemweb (ofertas.paginas, con las funciones de gemweb_extractor.paginas);
aqui solo se da formato.
"""

import io
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LOGO = Path(__file__).resolve().parents[1] / "assets" / "logo_geype.png"
TELEFONO_GEYPE = "(+34) 954 636 737"

MESES = ["Ene / Jan", "Feb / Feb", "Mar / Mar", "Abr / Apr", "May / May", "Jun / Jun",
         "Jul / Jul", "Ago / Aug", "Sep / Sep", "Oct / Oct", "Nov / Nov", "Dic / Dec"]
PERIODOS = ["P1", "P2", "P3", "P4", "P5", "P6"]
COLS_P = ["P1 kWh", "P2 kWh", "P3 kWh", "P4 kWh", "P5 kWh", "P6 kWh"]
# colores de las series del grafico (como en la plantilla): p1 rojo ... p6 verde
COLORES_P = ["FF0000", "7030A0", "FF6600", "FFFF00", "0000FF", "00A000"]

# tipologias de precio que se pueden pedir. El fijo, si se pide, es la binomica segun la
# tarifa (precio_fijo)
PRECIO_FIJO = "Precio fijo"
PRECIO_INDEXADO = ["PVPC", "Indexada Omie Pass Through + Omip (carga base)",
                   "Indexada Omie Pass Through + Omip (% consumo)",
                   "Indexada Omie Pass Pool + Omip (carga base)",
                   "Indexada Omie Pass Pool + Omip (% consumo)",
                   "Indexada Omie Pass Through", "Indexada Omie Pass Pool"]

# ------------------------------------------------------------------- gas natural
ELECTRICIDAD, GAS = "electricidad", "gas"
# lo que se escribe en «Producto / Commodity»
PRODUCTOS = {ELECTRICIDAD: "ELECTRICIDAD / ELECTRICITY", GAS: "GAS NATURAL / NATURAL GAS"}
# indexadas de gas: cada mercado, sin y con coberturas (fijar el precio de parte del consumo)
CON_COBERTURAS = " con coberturas"
PRECIO_INDEXADO_GAS = [f"Indexado {m}{c}" for m in ("MIBGAS (PVB)", "TTF", "Brent")
                       for c in ("", CON_COBERTURAS)]
# Impuesto sobre Hidrocarburos (Ley 38/1992): epigrafes 1.10.2 y 1.10.1 y exenciones del art. 51
IEH_REDUCIDO = "Tipo reducido 0,15 €/GJ (usos profesionales)"
IEH_GENERAL = "Tipo general 0,65 €/GJ"
IEH_EXENTO = "Exento"
TIPOS_IEH = [IEH_REDUCIDO, IEH_GENERAL, IEH_EXENTO]
IEH_INGLES = {IEH_REDUCIDO: "Reduced rate (professional use)", IEH_GENERAL: "General rate",
              IEH_EXENTO: "Exempt"}
# termino fijo del peaje (Circular 6/2020, art. 26): por cliente o por capacidad (Qd)
CUOTA_FIJA, VARIABLE = "Cuota fija", "Variable"
TERMINO_INGLES = {CUOTA_FIJA: "Cuota fija / Fixed fee",
                  VARIABLE: "Variable (por caudal) / Variable (by capacity)"}
# peajes RL por consumo anual (Circular 6/2020, art. 21): (consumo maximo en kWh, grupo)
ESCALONES_RL = [(5_000, 1), (15_000, 2), (50_000, 3), (300_000, 4), (1_500_000, 5),
                (5_000_000, 6), (15_000_000, 7), (50_000_000, 8), (150_000_000, 9),
                (500_000_000, 10), (float("inf"), 11)]
# telemedida obligatoria por encima de este consumo anual (Orden IET/2446/2013, art. 9)
UMBRAL_TELEMEDIDA = 5_000_000
# suministro interrumpible (Circular 6/2020, art. 24): consumo anual, diario y presion minimos
INTERRUMPIBLE_KWH, INTERRUMPIBLE_QD, INTERRUMPIBLE_BAR = 10_000_000, 26_000, 4
MESES_CORTOS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov",
                "dic"]
MESES_LARGOS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
                "septiembre", "octubre", "noviembre", "diciembre"]
# perfil de consumo: una factura mas larga que esto se reparte entre sus meses; los dias que
# cubren las facturas pueden alejarse de un año hasta MARGEN_DIAS sin avisar
FACTURA_MENSUAL_MAX_DIAS, DIAS_ANIO, MARGEN_DIAS = 45, 365, 15
AMARILLO, ROJO = "FFFF00", "FF0000"
# consumo que no sale de las facturas sino de la plantilla de ofertas.plantilla: se avisa bajo
# el perfil
CONSUMO_PLANTILLA = "Plantilla"
NOTA_CONSUMO = ("Consumo previsto facilitado por GE&PE / Expected consumption provided by "
                "GE&PE")

OSCURO, CLARO, NARANJA = "1F497D", "DAEEF3", "FFC000"
_borde = Side(style="thin", color="FFFFFF")
BORDE = Border(left=_borde, right=_borde, top=_borde, bottom=_borde)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ------------------------------------------------------------------- datos
def alias_de(nombre_suministro: str) -> str:
    """Todo lo que va tras el PRIMER '_': 'Empresa de prueba SA_TC4' -> 'TC4',
    'Empresa_Nave_2' -> 'Nave_2'. Sin '_', el nombre entero del suministro en Gemweb."""
    nombre = (nombre_suministro or "").strip()
    return nombre.split("_", 1)[1].strip() if "_" in nombre else nombre


def cif_limpio(cif: str) -> str:
    return re.sub(r"[\s.\-]", "", cif or "").upper()


def precio_fijo(tarifa: str) -> str:
    """Precio fijo, cuando se pide: binomica en DH2 para la 2.0TD, en DH6 para la 3.0TD y
    las 6.xTD."""
    return "Binómica en DH2" if str(tarifa or "").startswith("2.0") else "Binómica en DH6"


def consumo_anual(s: dict) -> float:
    """Suma del perfil de consumo (kWh) de un suministro (en gas, el perfil de gas)."""
    if s.get("perfil_gas"):
        return sum(p["kwh"] for p in s["perfil_gas"].values())
    return sum(sum(v) for v in (s.get("perfil") or {}).values())


def producto_de(tipo: str) -> str:
    """Tipo de suministro de Gemweb (tipus_submin) -> producto de la solicitud."""
    return GAS if str(tipo or "").strip().lower() == "gas" else ELECTRICIDAD


def grupo_rl(tarifa: str):
    """'RL.6', 'RL06', 'RLTB.6' -> 6 (None si no es un peaje RL)."""
    m = re.match(r"\s*RL[A-Z]*\.?0*(\d+)", str(tarifa or "").upper())
    return int(m.group(1)) if m else None


def tarifa_gas(tarifa: str) -> str:
    """Peaje como se escribe en las peticiones de gas: 'RL.6' -> 'RL06'."""
    n = grupo_rl(tarifa)
    return f"RL{n:02d}" if n else (tarifa or "")


def grupo_rl_de_consumo(consumo: float) -> int:
    return next(n for limite, n in ESCALONES_RL if consumo <= limite)


def reparto_mensual(f: dict) -> list:
    """Meses del perfil de una factura: [(año, mes, fraccion del consumo, dias)]. Una factura
    de hasta FACTURA_MENSUAL_MAX_DIAS va entera al mes con mas dias suyos (a igualdad, el
    ultimo); una mas larga (bimestral, trimestral...) se reparte entre sus meses segun los
    dias de cada uno. Sin fecha de inicio, al mes de su fecha de fin."""
    fin = f["Fin factura"]
    ini = f.get("Inicio factura") or fin
    total = (fin - ini).days + 1
    tramos, d = [], ini
    while d <= fin:
        fin_mes = date(d.year + d.month // 12, d.month % 12 + 1, 1) - timedelta(days=1)
        hasta = min(fin_mes, fin)
        tramos.append((d.year, d.month, (hasta - d).days + 1))
        d = hasta + timedelta(days=1)
    if total <= FACTURA_MENSUAL_MAX_DIAS:
        anio, mes, _ = max(tramos, key=lambda t: (t[2], t[0], t[1]))
        return [(anio, mes, 1.0, total if f.get("Inicio factura") else 0)]
    return [(anio, mes, dias / total, dias) for anio, mes, dias in tramos]


def perfil_gas(facturas: list) -> dict:
    """Facturas de gas -> {mes 1-12: {anio, kwh, dias, qd}}, con el reparto de
    reparto_mensual (como perfil_mensual); dias, los facturados en ese mes; qd, el Qd
    aplicado (kWh/dia) de la factura (el mayor si hay varias)."""
    perfil = {m: {"anio": None, "kwh": 0.0, "dias": 0, "qd": None} for m in range(1, 13)}
    for f in facturas:
        for anio, mes, fraccion, dias in reparto_mensual(f):
            p = perfil[mes]
            p["anio"] = max(p["anio"] or 0, anio)
            p["kwh"] += float(f.get("Total kWh") or 0) * fraccion
            p["dias"] += dias
            if f.get("_qd"):
                p["qd"] = max(p["qd"] or 0, float(f["_qd"]))
    return perfil


def tecnicos_por_defecto(consumo: float, qd=None) -> dict:
    """Aspectos tecnicos de un suministro de gas con lo que se sabe de Gemweb. La telemedida
    no esta en Gemweb (sale de Datadis, que no tiene gas): se supone si es obligatoria."""
    telemedida = consumo > UMBRAL_TELEMEDIDA
    return {"qd": qd, "telemedida": telemedida,
            "termino_fijo": VARIABLE if telemedida else CUOTA_FIJA, "presion": None,
            "cogeneracion": False, "planta_satelite": False, "dias": None,
            "interrumpible": False}


def reglas_tecnicas(t: dict, consumo: float) -> tuple:
    """Lo obligatorio: con mas de 5 GWh/año, telemedida y termino fijo por capacidad.
    -> (aspectos tecnicos corregidos, [lo que se ha corregido])."""
    t, cambios = dict(t), []
    if consumo > UMBRAL_TELEMEDIDA:
        if not t.get("telemedida"):
            t["telemedida"] = True
            cambios.append("telemedida (obligatoria con más de 5 GWh/año)")
        if t.get("termino_fijo") != VARIABLE:
            t["termino_fijo"] = VARIABLE
            cambios.append("término fijo variable (por capacidad, con telemedida)")
    return t, cambios


def ieh_valido(ieh: str, cogeneracion: bool) -> str:
    """Con cogeneracion no cabe el tipo reducido (epigrafe 1.10.2): se propone la exencion."""
    return IEH_EXENTO if cogeneracion and ieh == IEH_REDUCIDO else ieh


def _plazo_corto(c: dict) -> bool:
    """Alguna duracion del contrato de menos de un año desde la fecha de inicio."""
    if not c.get("inicio"):
        return False
    ini = datetime.strptime(c["inicio"], "%d/%m/%Y").date()
    return any((datetime.strptime(c[k], "%d/%m/%Y").date() - ini).days + 1 < 365
               for k in ("duracion", "duracion2", "duracion3") if c.get(k))


def avisos_gas(s: dict) -> list:
    """Avisos informativos de un suministro de gas (no impiden generar la solicitud)."""
    t, c = s.get("tecnicos") or {}, s.get("consideraciones") or {}
    consumo, qd = consumo_anual(s), t.get("qd")
    avisos = []
    grupo = grupo_rl(s.get("tarifa"))
    if grupo and consumo and grupo != grupo_rl_de_consumo(consumo):
        avisos.append("la tarifa %s no corresponde a su consumo anual (%s kWh: RL%02d)" % (
            tarifa_gas(s.get("tarifa")), f"{consumo:,.0f}".replace(",", "."),
            grupo_rl_de_consumo(consumo)))
    if t.get("telemedida") and qd and consumo / (qd * 365) > 0.8:
        avisos.append("factor de carga del %.0f %% (con telemedida no debe pasar del 80 %%)"
                      % (consumo / (qd * 365) * 100))
    if t.get("termino_fijo") == VARIABLE and not t.get("telemedida"):
        avisos.append("término fijo variable sin telemedida: hace falta un equipo que registre "
                      "el caudal diario")
    if not t.get("telemedida") and _plazo_corto(c):
        avisos.append("alguna duración es de menos de 12 meses y no tiene telemedida: la "
                      "capacidad de corto plazo la exige, así que se contratará capacidad anual")
    if t.get("interrumpible"):
        faltan = [texto for falta, texto in (
            (consumo <= INTERRUMPIBLE_KWH, "más de 10 GWh/año"),
            (not qd or qd <= INTERRUMPIBLE_QD, "más de 26.000 kWh/día"),
            (not t.get("presion") or t["presion"] <= INTERRUMPIBLE_BAR, "más de 4 bar"),
            (not t.get("telemedida"), "telemedida")) if falta]
        if faltan:
            avisos.append("interrumpible sin cumplir los requisitos: " + ", ".join(faltan))
    return avisos


def aviso_perfil(facturas: list) -> tuple:
    """(texto, grave): si los dias que cubren las facturas se alejan de un año mas de
    MARGEN_DIAS (sobran o faltan) o si hay facturas que se solapan mas de MARGEN_DIAS (p. ej.
    una complementaria del mismo periodo). ('', False) si todo cuadra, o si alguna factura no
    trae su fecha de inicio."""
    if not facturas or not all(f.get("Inicio factura") for f in facturas):
        return "", False
    dias = set()
    for f in facturas:
        dias.update(range(f["Inicio factura"].toordinal(), f["Fin factura"].toordinal() + 1))
    cubiertos = len(dias)
    # solapes solo entre facturas con consumo: una de 0 kWh del mismo periodo (complementaria
    # de otros conceptos, abono...) no suma nada dos veces
    con_consumo = [f for f in facturas if float(f.get("Total kWh", 1) or 0)]
    veces = {}
    for f in con_consumo:
        for d in range(f["Inicio factura"].toordinal(), f["Fin factura"].toordinal() + 1):
            veces[d] = veces.get(d, 0) + 1
    solape = sum(v - 1 for v in veces.values() if v > 1)
    pares = [(f, g) for n, f in enumerate(con_consumo) for g in con_consumo[n + 1:]
             if f["Inicio factura"] <= g["Fin factura"] and g["Inicio factura"] <= f["Fin factura"]]
    avisos = []
    if cubiertos > DIAS_ANIO + MARGEN_DIAS:
        avisos.append(f"las facturas cubren {cubiertos} días: sobran {cubiertos - DIAS_ANIO} "
                      "sobre un año y el consumo anual sale de más")
    elif cubiertos < DIAS_ANIO - MARGEN_DIAS:
        avisos.append(f"las facturas solo cubren {cubiertos} días: faltan "
                      f"{DIAS_ANIO - cubiertos} para un año (meses sin factura) y el consumo "
                      "anual sale de menos")
    if solape > MARGEN_DIAS:
        def periodo(f):
            kwh = f"{float(f.get('Total kWh') or 0):,.0f}".replace(",", ".")
            return (f"{f['Inicio factura']:%d/%m/%Y}-{f['Fin factura']:%d/%m/%Y} "
                    f"({kwh} kWh)" if "Total kWh" in f else
                    f"{f['Inicio factura']:%d/%m/%Y}-{f['Fin factura']:%d/%m/%Y}")
        detalle = "; ".join(f"{periodo(f)} y {periodo(g)}" for f, g in pares)
        avisos.append(f"hay facturas con consumo que se solapan {solape} días ({detalle}), p. "
                      "ej. una complementaria o rectificativa del mismo periodo: revisa en "
                      "Gemweb que el consumo no se cuente dos veces")
    return ("; además, ".join(avisos) + "." if avisos else ""), bool(avisos)


def perfil_mensual(facturas: list) -> tuple:
    """Facturas -> ({mes 1-12: [P1..P6]}, aviso). Cada factura mensual va al mes con mas dias
    suyos y las mas largas se reparten por dias (reparto_mensual); si un mes tiene varias, se
    suman. El aviso es el texto de aviso_perfil."""
    perfil = {m: [0.0] * 6 for m in range(1, 13)}
    for f in facturas:
        for _, mes, fraccion, _ in reparto_mensual(f):
            for n, col in enumerate(COLS_P):
                perfil[mes][n] += float(f.get(col) or 0) * fraccion
    return perfil, aviso_perfil(facturas)[0]


# ------------------------------------------------------------------- formato
def _celda(ws, ref, valor=None, oscura=False, negrita=None, formato=None):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.fill = PatternFill("solid", fgColor=OSCURO if oscura else CLARO)
    c.font = Font(name="Calibri", size=10, bold=oscura if negrita is None else negrita,
                  color="FFFFFF" if oscura else "000000")
    c.alignment = CENTRO
    c.border = BORDE
    if formato:
        c.number_format = formato
    return c


def _bloque(ws, rango, valor, oscura=False, formato=None):
    """Rango combinado con el mismo estilo en todas sus celdas (si no, el borde se corta)."""
    primera = rango.split(":")[0]
    if ":" in rango:
        ws.merge_cells(rango)
        for fila in ws[rango]:
            for c in fila:
                _celda(ws, c.coordinate, oscura=oscura, formato=formato)
    _celda(ws, primera, valor, oscura=oscura, formato=formato)


def _logo(ws, celda, alto_px=85):
    if LOGO.exists():
        img = Image(str(LOGO))
        img.width, img.height = int(img.width * alto_px / img.height), alto_px
        ws.add_image(img, celda)


def _pct(valor):
    return "" if valor is None else f"{valor:g} %".replace(".", ",")


def _si_no(ws, fila, texto, extra=()):
    """Fila doble «texto: Si / Yes», con datos opcionales a la derecha [(texto, valor)]."""
    _bloque(ws, f"B{fila}:D{fila + 1}", texto, True)
    _bloque(ws, f"E{fila}:E{fila + 1}", "Sí / Yes")
    for n, (etiqueta, valor) in enumerate(extra):
        _bloque(ws, f"F{fila + n}:H{fila + n}", etiqueta, True)
        _bloque(ws, f"I{fila + n}", valor)
    return fila + 3


def _ieh(ws, fila, c):
    """Fila doble del Impuesto sobre Hidrocarburos: tipo y, si lo hay, el % del consumo."""
    tipo = c.get("ieh") or IEH_REDUCIDO
    _bloque(ws, f"B{fila}:D{fila + 1}", "Impuesto sobre Hidrocarburos / Hydrocarbon tax", True)
    if c.get("pct_ieh") is None:
        _bloque(ws, f"E{fila}:I{fila + 1}", f"{tipo} / {IEH_INGLES.get(tipo, '')}")
    else:
        _bloque(ws, f"E{fila}:I{fila}", f"{tipo} / {IEH_INGLES.get(tipo, '')}")
        _bloque(ws, f"E{fila + 1}:H{fila + 1}", "% del consumo al que se aplica / % of "
                                                "consumption it applies to", True)
        _bloque(ws, f"I{fila + 1}", _pct(c["pct_ieh"]))
    return fila + 3


def _consideraciones(ws, cab, s, fila):
    """Bloque «Tipologia de ofertas solicitadas» bajo el grafico (empieza en `fila`).
    Devuelve la ultima fila usada."""
    c = s.get("consideraciones") or {}
    gas = cab.get("producto_tipo") == GAS
    _bloque(ws, f"B{fila}:I{fila}", "Tipología de ofertas solicitadas / Typology of requested "
                                    "offers", True)
    # precio fijo (si se pide, segun la tarifa; en gas no hay binomica): tres por fila
    # (D:E, F:G, H:I)
    fijo = [] if not c.get("fijo", True) else ["Sí / Yes"] if gas else [
        precio_fijo(s.get("tarifa"))]
    filas_fijo = max(1, -(-len(fijo) // 3))
    ini = fila + 1
    _bloque(ws, f"B{ini}:C{ini + filas_fijo - 1}", "Precio fijo / Fixed price", True)
    for n in range(filas_fijo * 3):
        f, col = ini + n // 3, "DFH"[n % 3]
        _bloque(ws, f"{col}{f}:{chr(ord(col) + 1)}{f}", fijo[n] if n < len(fijo) else "")
    # precio indexado: una opcion por fila, al menos tres filas
    indexado = [o.replace(CON_COBERTURAS, " (con coberturas / with hedging)")
                for o in c.get("indexado") or []]
    ini += filas_fijo
    filas_ind = max(3, len(indexado))
    _bloque(ws, f"B{ini}:C{ini + filas_ind - 1}", "Precio indexado / Indexed price", True)
    for n in range(filas_ind):
        ws.row_dimensions[ini + n].height = 27
        _bloque(ws, f"D{ini + n}:I{ini + n}", indexado[n] if n < len(indexado) else "")
    fila = ini + filas_ind + 1

    _bloque(ws, f"B{fila}:D{fila + 1}", "Tipo de contrato / Type of contract", True)
    _bloque(ws, f"E{fila}:E{fila + 1}",
            "De temporada / Seasonal" if c.get("temporada") else "Anual / Annual")
    fila += 3
    _bloque(ws, f"B{fila}:D{fila + 1}", "Fecha prevista de inicio del contrato / Expected date "
                                        "the contract starts", True)
    _bloque(ws, f"E{fila}:E{fila + 1}", c.get("inicio", ""))
    fila += 3
    plazos = [d for d in (c.get("duracion"), c.get("duracion2"), c.get("duracion3")) if d] or [""]
    _bloque(ws, f"B{fila}:C{fila + max(len(plazos), 2) - 1}",
            "Duración del contrato / contract duration", True)
    for n in range(max(len(plazos), 2)):
        _bloque(ws, f"D{fila + n}:E{fila + n}", "Hasta / Until", True)
        _bloque(ws, f"F{fila + n}", plazos[n] if n < len(plazos) else "")
    fila += max(len(plazos), 2) + 1
    # particularidades: solo salen si se aplican al suministro (el impuesto de gas, siempre)
    if gas:
        fila = _ieh(ws, fila, c)
        if c.get("gdo"):
            fila = _si_no(ws, fila, "Garantías de origen de gas renovable incluidas / "
                                    "Renewable gas guarantees of origin included")
    else:
        if c.get("exencion"):
            fila = _si_no(ws, fila, "Sujeto a exención de impuesto eléctrico / Subject to "
                                    "electricity tax exemption",
                          [("% consumo sujeto a exención / % exempt consumption",
                            _pct(c.get("pct_consumo_exento"))),
                           ("% tipo de exención / % exemption rate",
                            _pct(c.get("pct_tipo_exencion")))])
        if c.get("gdo"):
            fila = _si_no(ws, fila, "Garantías de origen (GdO) incluidas / Guarantees of "
                                    "origin included")
        if c.get("excedentes"):
            fila = _si_no(ws, fila, "Compensación de excedentes / Surplus compensation")
    _bloque(ws, f"B{fila}:D{fila + 1}", "Periodo para la recepción de las ofertas / Period "
                                        "for reception of offers", True)
    hora = cab.get("recepcion_hora", "")
    for n, (texto, valor) in enumerate((("Desde/Since:", cab.get("recepcion_desde", "")),
                                        ("Hasta/Until:", cab.get("recepcion_hasta", "")))):
        _bloque(ws, f"E{fila + n}", texto, True)
        celda = _celda(ws, f"F{fila + n}", valor, negrita=True)
        celda.fill = PatternFill("solid", fgColor=NARANJA)
    if hora:
        _bloque(ws, f"G{fila + 1}:H{fila + 1}", f"antes de las {hora} h / before {hora}")
        for col in "GH":
            ws[f"{col}{fila + 1}"].fill = PatternFill("solid", fgColor=NARANJA)
            ws[f"{col}{fila + 1}"].font = Font(name="Calibri", size=10, bold=True)
    return fila + 1


# ------------------------------------------------------------------- hojas
def _potencias_y_perfil(ws, s):
    """Electricidad: tarifa, potencias, perfil por periodos P1-P6 y su grafico."""
    _bloque(ws, "B18", "Tarifa / Tariff", True)
    _bloque(ws, "C18", s.get("tarifa", ""))
    _bloque(ws, "D19:I19", "Potencias a contratar / Powers to contract (kW)", True)
    for n, col in enumerate("DEFGHI"):
        _bloque(ws, f"{col}20", PERIODOS[n], True)
        pot = (s.get("potencias") or [None] * 6)[n]
        _bloque(ws, f"{col}21", pot, formato="General")

    _bloque(ws, "B23:I23", "Perfil de consumo previsto según / Expected consumption profile "
                           "according to (kWh)", True)
    _bloque(ws, "B24", "Mes / Month", True)
    for n, col in enumerate("CDEFGH"):
        _bloque(ws, f"{col}24", PERIODOS[n], True)
    _bloque(ws, "I24", "Total / Overall", True)
    perfil = s.get("perfil") or {m: [0.0] * 6 for m in range(1, 13)}
    for m in range(1, 13):
        fila = 24 + m
        _bloque(ws, f"B{fila}", MESES[m - 1], True)
        for n, col in enumerate("CDEFGH"):
            _bloque(ws, f"{col}{fila}", perfil[m][n], formato="#,##0")
        _bloque(ws, f"I{fila}", f"=SUM(C{fila}:H{fila})", True, formato="#,##0")
    _bloque(ws, "B37", "Total / Overall", True)
    _bloque(ws, "B38", "%", True)
    for col in "CDEFGHI":
        _bloque(ws, f"{col}37", f"=SUM({col}25:{col}36)", True, formato="#,##0")
    for col in "CDEFGH":
        _bloque(ws, f"{col}38", f"=IF($I$37=0,0,{col}37/$I$37*100)", True, formato="#,##0.00")
    _bloque(ws, "I38", None, True)
    _nota_consumo(ws, s, "B39")

    grafico = BarChart()
    grafico.type, grafico.grouping, grafico.overlap = "col", "stacked", 100
    grafico.y_axis.title = "kWh"
    grafico.y_axis.title.overlay = False            # el titulo no se monta sobre los numeros
    grafico.y_axis.numFmt = "#,##0"
    grafico.y_axis.majorGridlines = None
    grafico.legend.position = "t"
    # sin esto, Excel moderno oculta los ejes de los graficos de openpyxl
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False
    grafico.width, grafico.height = 23.5, 8.6
    datos = Reference(ws, min_col=3, max_col=8, min_row=24, max_row=36)
    grafico.add_data(datos, titles_from_data=True)
    grafico.set_categories(Reference(ws, min_col=2, min_row=25, max_row=36))
    for serie, color, nombre in zip(grafico.series, COLORES_P, PERIODOS):
        serie.graphicalProperties.solidFill = color
        serie.graphicalProperties.line.solidFill = color
    grafico.gapWidth = 60
    ws.add_chart(grafico, "B41")


def _si_no_texto(valor) -> str:
    return "SÍ" if valor else "NO"


def _o_guion(valor):
    return "-" if valor is None or valor == "" else valor


def _aspectos_tecnicos(ws, s, fila=18):
    """Gas: bloque «Aspectos tecnicos» (filas 18-24), con la disposicion de las peticiones de
    gas de GE&PE. El consumo anual es el total del perfil (E40)."""
    t = s.get("tecnicos") or {}
    _bloque(ws, f"B{fila}:I{fila}", "Aspectos técnicos / Technical aspects", True)
    for f in range(fila + 1, fila + 7):                # etiquetas bilingües de hasta 3 lineas
        ws.row_dimensions[f].height = 42 if f < fila + 6 else 27
    f = fila + 1
    _bloque(ws, f"B{f}:C{f}", "Tarifa a contratar / Tariff", True)
    _bloque(ws, f"D{f}", tarifa_gas(s.get("tarifa")))
    _bloque(ws, f"E{f}:F{f}", "Presión de suministro (bar) / Supply pressure (bar)", True)
    _bloque(ws, f"G{f}", _o_guion(t.get("presion")), formato="General")
    _bloque(ws, f"H{f}", "Consumo anual / Annual consumption (kWh)", True)
    _bloque(ws, f"I{f}", "=E40", formato="#,##0")
    _bloque(ws, f"B{f + 1}:C{f + 1}", "Cogeneración (SÍ/NO) / Cogeneration (Y/N)", True)
    _bloque(ws, f"D{f + 1}", _si_no_texto(t.get("cogeneracion")))
    _bloque(ws, f"B{f + 2}:C{f + 2}", "Alquiler de planta satélite de regasificación (SÍ/NO) / "
                                      "Satellite regasification plant rental (Y/N)", True)
    _bloque(ws, f"D{f + 2}", _si_no_texto(t.get("planta_satelite")))
    _bloque(ws, f"E{f + 2}:H{f + 2}", "Días de utilización a la semana / Days of use per week",
            True)
    _bloque(ws, f"I{f + 2}", _o_guion(t.get("dias")), formato="General")
    _bloque(ws, f"B{f + 3}:C{f + 3}", "Qd a contratar (kWh/día) / Daily capacity to contract "
                                      "(kWh/day)", True)
    _bloque(ws, f"D{f + 3}", _o_guion(t.get("qd")), formato="#,##0")
    _bloque(ws, f"B{f + 4}:C{f + 4}", "Suministro interrumpible (SÍ/NO) / Interruptible supply "
                                      "(Y/N)", True)
    _bloque(ws, f"D{f + 4}", _si_no_texto(t.get("interrumpible")))
    _bloque(ws, f"E{f + 4}:H{f + 4}", "Telemedida (SÍ/NO) / Telemetering (Y/N)", True)
    _bloque(ws, f"I{f + 4}", _si_no_texto(t.get("telemedida")))
    _bloque(ws, f"B{f + 5}:C{f + 5}", "Término fijo / Fixed term", True)
    termino = t.get("termino_fijo") or CUOTA_FIJA
    _bloque(ws, f"D{f + 5}:E{f + 5}", TERMINO_INGLES.get(termino, termino))
    for col in "DE":                                # destacado como en la plantilla de gas
        ws[f"{col}{f + 5}"].fill = PatternFill("solid", fgColor=AMARILLO)
        ws[f"{col}{f + 5}"].font = Font(name="Calibri", size=10, bold=True, color=ROJO)


def _nota_consumo(ws, s, celda):
    if s.get("origen_consumo") == CONSUMO_PLANTILLA:
        ws[celda] = NOTA_CONSUMO
        ws[celda].font = Font(italic=True, size=9, color=OSCURO)


def _perfil_gas(ws, s, fila=26):
    """Gas: perfil de consumo de 12 meses (Qd medio, energia y Qd aplicado; total en E40) y
    grafico de la energia."""
    perfil = s.get("perfil_gas") or {m: {"anio": None, "kwh": 0.0, "dias": 0, "qd": None}
                                     for m in range(1, 13)}
    _bloque(ws, f"B{fila}:I{fila}", "Perfil de consumo / Consumption profile (kWh)", True)
    cab = fila + 1
    ws.row_dimensions[cab].height = 27
    _bloque(ws, f"B{cab}", "Mes / Month", True)
    _bloque(ws, f"C{cab}:D{cab}", "Qd media (kWh/día) / Average Qd (kWh/day)", True)
    _bloque(ws, f"E{cab}:F{cab}", "Energía consumida (kWh) / Consumed energy (kWh)", True)
    _bloque(ws, f"G{cab}:I{cab}", "Qd aplicada (kWh/día) / Applied Qd (kWh/day)", True)
    for m in range(1, 13):
        p, f = perfil[m], cab + m
        anio = p.get("anio")
        _bloque(ws, f"B{f}", f"{MESES_CORTOS[m - 1]}-{anio % 100:02d}" if anio
                else MESES[m - 1], True)
        media = p["kwh"] / p["dias"] if p.get("dias") else None
        _bloque(ws, f"C{f}:D{f}", _o_guion(round(media) if media is not None else None),
                formato="#,##0")
        _bloque(ws, f"E{f}:F{f}", p["kwh"], formato="#,##0")
        _bloque(ws, f"G{f}:I{f}", _o_guion(p.get("qd")), formato="#,##0")
    total = cab + 13
    _bloque(ws, f"B{total}", "Total / Overall", True)
    _bloque(ws, f"C{total}:D{total}", None, True)
    _bloque(ws, f"E{total}:F{total}", f"=SUM(E{cab + 1}:E{cab + 12})", True, formato="#,##0")
    _bloque(ws, f"G{total}:I{total}", None, True)
    _nota_consumo(ws, s, f"B{total + 1}")

    grafico = BarChart()
    grafico.type = "col"
    grafico.y_axis.title = "kWh"
    grafico.y_axis.title.overlay = False
    grafico.y_axis.numFmt = "#,##0"
    grafico.y_axis.majorGridlines = None
    grafico.legend = None
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False
    grafico.width, grafico.height = 23.5, 8.6
    grafico.add_data(Reference(ws, min_col=5, min_row=cab, max_row=cab + 12),
                     titles_from_data=True)
    grafico.set_categories(Reference(ws, min_col=2, min_row=cab + 1, max_row=cab + 12))
    grafico.series[0].graphicalProperties.solidFill = OSCURO
    grafico.series[0].graphicalProperties.line.solidFill = OSCURO
    grafico.gapWidth = 60
    ws.add_chart(grafico, f"B{total + 2}")

def _hoja_suministro(wb, nombre_hoja, cab, s):
    ws = wb.create_sheet(nombre_hoja)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3.43
    ws.column_dimensions["B"].width = 12.29
    for col in "CDEFGHI":
        ws.column_dimensions[col].width = 15.7
    ws.column_dimensions["J"].width = 3.43
    for fila in range(1, 60):
        ws.row_dimensions[fila].height = 13.5
    _logo(ws, "E1")

    _bloque(ws, "B6", "Ref.", True)
    _bloque(ws, "C6:D6", cab.get("ref", ""))
    _bloque(ws, "H6", "Fecha / Date", True)
    _bloque(ws, "I6", cab.get("fecha", ""))
    _bloque(ws, "B8:C8", "Gestor de Cuenta / KAM", True)
    _bloque(ws, "D8:F8", cab.get("gestor", ""))
    _bloque(ws, "G8", "Tlfno./ phone", True)
    _bloque(ws, "H8:I8", cab.get("telefono", ""))
    _bloque(ws, "G9", "email", True)
    _bloque(ws, "H9:I9", cab.get("email", ""))
    _bloque(ws, "B10:C10", "Producto / Commodity", True)
    _bloque(ws, "D10:F10", cab.get("producto", ""))

    _bloque(ws, "B12:D12", "Razón Social / Company name", True)
    _bloque(ws, "E12:I12", s.get("razon", ""))
    _bloque(ws, "B13:C13", "CIF / NIF / ID number", True)
    _bloque(ws, "D13", s.get("cif", ""))
    _bloque(ws, "B14:C14", "CUPS / Supply point code", True)
    _bloque(ws, "D14:F14", s.get("cups", ""))
    _bloque(ws, "G14", "Alias", True)
    _bloque(ws, "H14:I14", s.get("alias", ""))
    _bloque(ws, "B15:C16", "Dirección de suministro / Supply address", True)
    _bloque(ws, "D15:I16", s.get("direccion", ""))

    if cab.get("producto_tipo") == GAS:
        _aspectos_tecnicos(ws, s)
        _perfil_gas(ws, s)
        ultima = _consideraciones(ws, cab, s, 62)
    else:
        _potencias_y_perfil(ws, s)
        ultima = _consideraciones(ws, cab, s, 61)
    ws.print_area = f"A1:J{ultima + 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


CAMPOS_RESUMEN = [("grupo", "Cliente/Grupo", 28), ("alias", "Alias", 10), ("cups", "CUPS", 25),
                  ("razon", "Razón social", 32), ("direccion", "Dirección de suministro", 45),
                  ("cif", "CIF", 12), ("tarifa", "Tarifa de acceso", 10),
                  ("fecha_inicio", "Fecha prevista de inicio del contrato", 13)]


def _hoja_resumen_gas(wb, suministros, nombres_hojas):
    """Gas: los mismos datos del suministro, sus aspectos tecnicos principales (Qd,
    telemedida, termino fijo) y el consumo anual, con la fila TOTAL bajo el consumo."""
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    campos = CAMPOS_RESUMEN
    tec0 = len(campos) + 1                  # primera columna de aspectos tecnicos
    total = tec0 + 3                        # consumo anual

    def letra(col):
        return ws.cell(1, col).column_letter

    for n, (_, texto, ancho) in enumerate(campos, start=1):
        ws.column_dimensions[letra(n)].width = ancho
        _bloque(ws, f"{letra(n)}2:{letra(n)}3", texto, True)
    for col, ancho in zip(range(tec0, total + 1), (12, 11, 14, 14)):
        ws.column_dimensions[letra(col)].width = ancho
    _bloque(ws, f"{letra(tec0)}2:{letra(tec0 + 2)}2", "ASPECTOS TÉCNICOS", True)
    for n, texto in enumerate(("Qd (kWh/día)", "Telemedida", "Término fijo")):
        _bloque(ws, ws.cell(3, tec0 + n).coordinate, texto, True)
    _bloque(ws, f"{letra(total)}2:{letra(total)}3", "Consumo anual (kWh)", True)
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24
    for fila, (s, hoja) in enumerate(zip(suministros, nombres_hojas), start=4):
        ws.row_dimensions[fila].height = 27
        t = s.get("tecnicos") or {}
        for n, (clave, _, _) in enumerate(campos, start=1):
            valor = tarifa_gas(s.get("tarifa")) if clave == "tarifa" else s.get(clave, "")
            _bloque(ws, ws.cell(fila, n).coordinate, valor)
        for n, valor in enumerate((_o_guion(t.get("qd")), _si_no_texto(t.get("telemedida")),
                                   t.get("termino_fijo") or CUOTA_FIJA)):
            _bloque(ws, ws.cell(fila, tec0 + n).coordinate, valor, formato="#,##0")
        _bloque(ws, ws.cell(fila, total).coordinate, f"='{hoja}'!E40", True, formato="#,##0")
    primera, ultima = 4, 3 + len(suministros)
    fila = ultima + 1
    ws.row_dimensions[fila].height = 20
    c = letra(total)
    _bloque(ws, f"{letra(total - 1)}{fila}", "TOTAL", True)
    _bloque(ws, f"{c}{fila}", f"=SUM({c}{primera}:{c}{ultima})" if suministros else 0, True,
            formato="#,##0")
    ws.freeze_panes = "D4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _hoja_resumen(wb, suministros, nombres_hojas):
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    campos = CAMPOS_RESUMEN
    pot0 = len(campos) + 1                  # primera columna de potencias
    con0 = pot0 + 6                         # primera columna de consumos
    total = con0 + 6                        # consumo anual

    def letra(col):
        return ws.cell(1, col).column_letter

    for n, (_, _, ancho) in enumerate(campos, start=1):
        ws.column_dimensions[letra(n)].width = ancho
    for col in range(pot0, total):
        ws.column_dimensions[letra(col)].width = 8
    ws.column_dimensions[letra(total)].width = 14
    for n, (_, texto, _) in enumerate(campos, start=1):
        _bloque(ws, f"{letra(n)}2:{letra(n)}3", texto, True)
    _bloque(ws, f"{letra(pot0)}2:{letra(pot0 + 5)}2", "POTENCIAS A CONTRATAR (kW)", True)
    _bloque(ws, f"{letra(con0)}2:{letra(total)}2", "PREVISIÓN DE CONSUMO ANUAL", True)
    for n, p in enumerate(PERIODOS):
        _bloque(ws, ws.cell(3, pot0 + n).coordinate, p, True)
        _bloque(ws, ws.cell(3, con0 + n).coordinate, p, True)
    _bloque(ws, ws.cell(3, total).coordinate, "Consumo anual  (kWh)", True)
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24
    for fila, (s, hoja) in enumerate(zip(suministros, nombres_hojas), start=4):
        ws.row_dimensions[fila].height = 27
        for n, (clave, _, _) in enumerate(campos, start=1):
            _bloque(ws, ws.cell(fila, n).coordinate, s.get(clave, ""))
        for n in range(6):
            pot = (s.get("potencias") or [None] * 6)[n]
            _bloque(ws, ws.cell(fila, pot0 + n).coordinate, pot, formato="General")
            col = "CDEFGH"[n]
            _bloque(ws, ws.cell(fila, con0 + n).coordinate, f"='{hoja}'!{col}37",
                    formato="#,##0")
        _bloque(ws, ws.cell(fila, total).coordinate, f"='{hoja}'!I37", True, formato="#,##0")
    # fila TOTAL: solo el consumo anual de todos los suministros, con la etiqueta a su izquierda
    primera, ultima = 4, 3 + len(suministros)
    fila = ultima + 1
    ws.row_dimensions[fila].height = 20
    c = letra(total)
    _bloque(ws, f"{letra(total - 1)}{fila}", "TOTAL", True)
    _bloque(ws, f"{c}{fila}", f"=SUM({c}{primera}:{c}{ultima})" if suministros else 0, True,
            formato="#,##0")
    ws.freeze_panes = "D4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def nombres_de_hojas(suministros: list) -> list:
    """'1. TC4', '2. TC3'... (una por suministro; sin alias, el final del CUPS)."""
    usados, nombres = set(), []
    for n, s in enumerate(suministros, start=1):
        base = re.sub(r"[\[\]:*?/\\']", "", s.get("alias") or s.get("cups", "")[-8:])
        nombre = f"{n}. {base}"[:31]
        while nombre in usados:
            nombre = nombre[:28] + "_" + str(n)
        usados.add(nombre)
        nombres.append(nombre)
    return nombres


def excel_solicitud(cabecera: dict, suministros: list) -> bytes:
    """cabecera: ref, fecha, gestor, telefono, email, producto y producto_tipo (ELECTRICIDAD
    o GAS). suministros: dicts con alias, cups, razon, direccion, cif, tarifa, fecha_inicio y
    consideraciones (ver ofertas.paginas.consideraciones_de); en electricidad, potencias [6] y
    perfil {mes: [6]}; en gas, tecnicos (tecnicos_por_defecto) y perfil_gas (perfil_gas)."""
    wb = Workbook()
    nombres = nombres_de_hojas(suministros)
    if cabecera.get("producto_tipo") == GAS:
        _hoja_resumen_gas(wb, suministros, nombres)
    else:
        _hoja_resumen(wb, suministros, nombres)
    for s, nombre in zip(suministros, nombres):
        _hoja_suministro(wb, nombre, cabecera, s)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
