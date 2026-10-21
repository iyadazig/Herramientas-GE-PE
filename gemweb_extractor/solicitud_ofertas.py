# -*- coding: utf-8 -*-
"""
SOLICITUD DE OFERTAS (Excel con el formato de GE&PE)
====================================================
Hoja «Resumen» con un suministro por fila y una hoja por suministro con la cabecera de la
solicitud, los datos del suministro, las potencias de la ultima factura y el perfil de
consumo de los ultimos 12 meses por periodos (cada factura en el mes de su fecha de fin),
con su grafico de columnas apiladas.

Los datos salen de la descarga masiva de consumos por periodos (gemweb_extractor.paginas);
aqui solo se da formato.
"""

import io
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LOGO = Path(__file__).resolve().parents[1] / "assets" / "logo_geype.png"
TELEFONO_GEYPE = "(+34) 954 636 737"

MESES = ["Ene / Jan", "Feb / Feb", "Mar / Mar", "Abr / Apr", "May / May", "Jun / Jun",
         "Jul / Jul", "Ago / Aug", "Sep / Sep", "Oct / Oct", "Nov / Nov", "Dic / Dec"]
PERIODOS = ["P1", "P2", "P3", "P4", "P5", "P6"]
COLS_P = ["P1 kWh", "P2 kWh", "P3 kWh", "P4 kWh", "P5 kWh", "P6 kWh"]
# colores de las series del grafico (como en la plantilla): p1 rojo ... p6 verde
COLORES_P = ["FF0000", "7030A0", "FF6600", "FFFF00", "0000FF", "00A000"]

OSCURO, CLARO = "1F497D", "DAEEF3"
_borde = Side(style="thin", color="FFFFFF")
BORDE = Border(left=_borde, right=_borde, top=_borde, bottom=_borde)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ------------------------------------------------------------------- datos
def alias_de(nombre_suministro: str) -> str:
    """Todo lo que va tras el PRIMER '_': 'Empresa de prueba SA_TC4' -> 'TC4',
    'Empresa_Nave_2' -> 'Nave_2'."""
    nombre = (nombre_suministro or "").strip()
    return nombre.split("_", 1)[1].strip() if "_" in nombre else ""


def cif_limpio(cif: str) -> str:
    return re.sub(r"[\s.\-]", "", cif or "").upper()


def discriminacion(tarifa: str) -> str:
    """Perfil de discriminacion horaria: DH3 para la 2.0TD, DH6 para el resto."""
    return "DH3" if str(tarifa or "").startswith("2.0") else "DH6"


def perfil_mensual(facturas: list) -> tuple:
    """Facturas -> ({mes 1-12: [P1..P6]}, aviso). Cada factura va al mes de su fecha de fin;
    si un mes tiene varias, se suman. El aviso dice si se mezclan años en un mismo mes."""
    perfil = {m: [0.0] * 6 for m in range(1, 13)}
    anios = {}
    for f in facturas:
        mes = f["Fin factura"].month
        anios.setdefault(mes, set()).add(f["Fin factura"].year)
        for n, col in enumerate(COLS_P):
            perfil[mes][n] += float(f.get(col) or 0)
    repetidos = sorted(m for m, a in anios.items() if len(a) > 1)
    aviso = ("Se suman facturas de años distintos en: %s" % ", ".join(MESES[m - 1][:3]
                                                                         for m in repetidos)
             if repetidos else "")
    return perfil, aviso


# ------------------------------------------------------------------- formato
def _celda(ws, ref, valor=None, oscura=False, negrita=None, formato=None):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.fill = PatternFill("solid", fgColor=OSCURO if oscura else CLARO)
    c.font = Font(name="Calibri", size=10, bold=oscura if negrita is None else negrita,
                  color="FFFFFF" if oscura else "000000")
    c.alignment = CENTRO
    c.border = BORDE
    if formato:
        c.number_format = formato
    return c


def _bloque(ws, rango, valor, oscura=False, formato=None):
    """Rango combinado con el mismo estilo en todas sus celdas (si no, el borde se corta)."""
    primera = rango.split(":")[0]
    if ":" in rango:
        ws.merge_cells(rango)
        for fila in ws[rango]:
            for c in fila:
                _celda(ws, c.coordinate, oscura=oscura, formato=formato)
    _celda(ws, primera, valor, oscura=oscura, formato=formato)


def _logo(ws, celda, alto_px=85):
    if LOGO.exists():
        img = Image(str(LOGO))
        img.width, img.height = int(img.width * alto_px / img.height), alto_px
        ws.add_image(img, celda)


# ------------------------------------------------------------------- hojas
def _hoja_suministro(wb, nombre_hoja, cab, s):
    ws = wb.create_sheet(nombre_hoja)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3.43
    ws.column_dimensions["B"].width = 12.29
    for col in "CDEFGHI":
        ws.column_dimensions[col].width = 15.7
    ws.column_dimensions["J"].width = 3.43
    for fila in range(1, 60):
        ws.row_dimensions[fila].height = 13.5
    _logo(ws, "E1")

    _bloque(ws, "B6", "Ref.", True)
    _bloque(ws, "C6:D6", cab.get("ref", ""))
    _bloque(ws, "H6", "Fecha / Date", True)
    _bloque(ws, "I6", cab.get("fecha", ""))
    _bloque(ws, "B8:C8", "Gestor de Cuenta / KAM", True)
    _bloque(ws, "D8:F8", cab.get("gestor", ""))
    _bloque(ws, "G8", "Tlfno./ phone", True)
    _bloque(ws, "H8:I8", cab.get("telefono", ""))
    _bloque(ws, "G9", "email", True)
    _bloque(ws, "H9:I9", cab.get("email", ""))
    _bloque(ws, "B10:C10", "Producto / Commodity", True)
    _bloque(ws, "D10:F10", cab.get("producto", ""))

    _bloque(ws, "B12:D12", "Razón Social / Company name", True)
    _bloque(ws, "E12:I12", s.get("razon", ""))
    _bloque(ws, "B13:C13", "CIF / NIF / ID number", True)
    _bloque(ws, "D13", s.get("cif", ""))
    _bloque(ws, "B14:C14", "CUPS / Supply point code", True)
    _bloque(ws, "D14:F14", s.get("cups", ""))
    _bloque(ws, "G14", "Alias", True)
    _bloque(ws, "H14:I14", s.get("alias", ""))
    _bloque(ws, "B15:C16", "Dirección de suministro / Supply address", True)
    _bloque(ws, "D15:I16", s.get("direccion", ""))

    _bloque(ws, "B18", "Tarifa / Tariff", True)
    _bloque(ws, "C18", s.get("tarifa", ""))
    _bloque(ws, "D19:I19", "Potencias a contratar / Powers to contract (kW)", True)
    for n, col in enumerate("DEFGHI"):
        _bloque(ws, f"{col}20", PERIODOS[n], True)
        pot = (s.get("potencias") or [None] * 6)[n]
        _bloque(ws, f"{col}21", pot, formato="General")

    _bloque(ws, "B23:G23", "Perfil de consumo previsto según / Expected consumption profile "
                           "according to (kWh)", True)
    _bloque(ws, "H23:I23", discriminacion(s.get("tarifa")))
    _bloque(ws, "B24", "Mes / Month", True)
    for n, col in enumerate("CDEFGH"):
        _bloque(ws, f"{col}24", PERIODOS[n], True)
    _bloque(ws, "I24", "Total / Overall", True)
    perfil = s.get("perfil") or {m: [0.0] * 6 for m in range(1, 13)}
    for m in range(1, 13):
        fila = 24 + m
        _bloque(ws, f"B{fila}", MESES[m - 1], True)
        for n, col in enumerate("CDEFGH"):
            _bloque(ws, f"{col}{fila}", perfil[m][n], formato="#,##0")
        _bloque(ws, f"I{fila}", f"=SUM(C{fila}:H{fila})", True, formato="#,##0")
    _bloque(ws, "B37", "Total / Overall", True)
    _bloque(ws, "B38", "%", True)
    for col in "CDEFGHI":
        _bloque(ws, f"{col}37", f"=SUM({col}25:{col}36)", True, formato="#,##0")
    for col in "CDEFGH":
        _bloque(ws, f"{col}38", f"=IF($I$37=0,0,{col}37/$I$37*100)", True, formato="#,##0.00")
    _bloque(ws, "I38", None, True)

    grafico = BarChart()
    grafico.type, grafico.grouping, grafico.overlap = "col", "stacked", 100
    grafico.y_axis.title = "kWh"
    grafico.y_axis.title.overlay = False            # el titulo no se monta sobre los numeros
    grafico.y_axis.numFmt = "#,##0"
    grafico.y_axis.majorGridlines = None
    grafico.legend.position = "t"
    # sin esto, Excel moderno oculta los ejes de los graficos de openpyxl
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False
    grafico.width, grafico.height = 23.5, 8.6
    datos = Reference(ws, min_col=3, max_col=8, min_row=24, max_row=36)
    grafico.add_data(datos, titles_from_data=True)
    grafico.set_categories(Reference(ws, min_col=2, min_row=25, max_row=36))
    for serie, color, nombre in zip(grafico.series, COLORES_P, PERIODOS):
        serie.graphicalProperties.solidFill = color
        serie.graphicalProperties.line.solidFill = color
    grafico.gapWidth = 60
    ws.add_chart(grafico, "B41")
    ws.print_area = "A1:J60"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _hoja_resumen(wb, suministros, nombres_hojas):
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    campos = [("grupo", "Cliente/Grupo", 28), ("alias", "Alias", 10), ("cups", "CUPS", 25),
              ("razon", "Razón social", 32), ("direccion", "Dirección de suministro", 45),
              ("cif", "CIF", 12), ("tarifa", "Tarifa de acceso", 10),
              ("fecha_fin", "Fecha fin contrato actual", 13)]
    pot0 = len(campos) + 1                  # primera columna de potencias
    con0 = pot0 + 6                         # primera columna de consumos
    total = con0 + 6                        # consumo anual

    def letra(col):
        return ws.cell(1, col).column_letter

    for n, (_, _, ancho) in enumerate(campos, start=1):
        ws.column_dimensions[letra(n)].width = ancho
    for col in range(pot0, total):
        ws.column_dimensions[letra(col)].width = 8
    ws.column_dimensions[letra(total)].width = 14
    for n, (_, texto, _) in enumerate(campos, start=1):
        _bloque(ws, f"{letra(n)}2:{letra(n)}3", texto, True)
    _bloque(ws, f"{letra(pot0)}2:{letra(pot0 + 5)}2", "POTENCIAS A CONTRATAR (kW)", True)
    _bloque(ws, f"{letra(con0)}2:{letra(total)}2", "PREVISIÓN DE CONSUMO ANUAL", True)
    for n, p in enumerate(PERIODOS):
        _bloque(ws, ws.cell(3, pot0 + n).coordinate, p, True)
        _bloque(ws, ws.cell(3, con0 + n).coordinate, p, True)
    _bloque(ws, ws.cell(3, total).coordinate, "Consumo anual  (kWh)", True)
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24
    for fila, (s, hoja) in enumerate(zip(suministros, nombres_hojas), start=4):
        ws.row_dimensions[fila].height = 27
        for n, (clave, _, _) in enumerate(campos, start=1):
            _bloque(ws, ws.cell(fila, n).coordinate, s.get(clave, ""))
        for n in range(6):
            pot = (s.get("potencias") or [None] * 6)[n]
            _bloque(ws, ws.cell(fila, pot0 + n).coordinate, pot, formato="General")
            col = "CDEFGH"[n]
            _bloque(ws, ws.cell(fila, con0 + n).coordinate, f"='{hoja}'!{col}37",
                    formato="#,##0")
        _bloque(ws, ws.cell(fila, total).coordinate, f"='{hoja}'!I37", True, formato="#,##0")
    # fila TOTAL: suma de los consumos de todos los suministros (las potencias no se suman)
    primera, ultima = 4, 3 + len(suministros)
    fila = ultima + 1
    ws.row_dimensions[fila].height = 20
    _bloque(ws, f"A{fila}:{letra(con0 - 1)}{fila}", "TOTAL", True)
    for col in range(con0, total + 1):
        c = letra(col)
        _bloque(ws, f"{c}{fila}", f"=SUM({c}{primera}:{c}{ultima})" if suministros else 0, True,
                formato="#,##0")
    ws.freeze_panes = "D4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def nombres_de_hojas(suministros: list) -> list:
    """'1. TC4', '2. TC3'... (una por suministro; sin alias, el final del CUPS)."""
    usados, nombres = set(), []
    for n, s in enumerate(suministros, start=1):
        base = re.sub(r"[\[\]:*?/\\']", "", s.get("alias") or s.get("cups", "")[-8:])
        nombre = f"{n}. {base}"[:31]
        while nombre in usados:
            nombre = nombre[:28] + "_" + str(n)
        usados.add(nombre)
        nombres.append(nombre)
    return nombres


def excel_solicitud(cabecera: dict, suministros: list) -> bytes:
    """cabecera: ref, fecha, gestor, telefono, email, producto. suministros: dicts con alias,
    cups, razon, direccion, cif, tarifa, fecha_fin, potencias [6] y perfil {mes: [6]}."""
    wb = Workbook()
    nombres = nombres_de_hojas(suministros)
    _hoja_resumen(wb, suministros, nombres)
    for s, nombre in zip(suministros, nombres):
        _hoja_suministro(wb, nombre, cabecera, s)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
